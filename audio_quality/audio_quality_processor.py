from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import uuid
from dataclasses import asdict, dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable

from .text_compat import compact_text, repair_multiline_text, repair_text


class AudioQualityError(ValueError):
    """Raised when an audio task or quality result cannot be accepted."""


@dataclass(frozen=True)
class IncomingAudioMessage:
    message_id: str
    group_id: str
    sender_user_id: str
    sender_name: str
    text: str
    attachment_id: str
    filename: str
    file_size: int | None = None
    duration_seconds: float | None = None
    attachment_hash: str | None = None


@dataclass(frozen=True)
class AudioSubmission:
    employee_name: str
    order_id: str | None
    accepted_number: str | None = None


@dataclass(frozen=True)
class TranscriptSegment:
    start_seconds: float
    end_seconds: float
    speaker: str
    text: str


@dataclass(frozen=True)
class RuleHit:
    rule_id: str
    start_seconds: float
    quote: str
    description: str
    deduction: int


@dataclass(frozen=True)
class QualityIssue:
    rule_id: str
    start_seconds: float
    quote: str
    reason: str
    suggestion: str
    confidence: float


@dataclass(frozen=True)
class QualityResult:
    employee_speaker: str
    speaker_confidence: float
    dimension_scores: dict[str, int]
    score: int
    qualified: bool
    summary: str
    analysis: str
    optimization: str
    issues: list[QualityIssue]
    needs_human_review: bool


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def normalize_text(value: Any) -> str:
    return repair_text(value)


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig") as handle:
        value = json.load(handle)
    if not isinstance(value, dict):
        raise AudioQualityError(f"JSON 根节点必须是对象：{path}")
    return value


def extract_accepted_number(*values: str) -> str | None:
    joined = " ".join(normalize_text(value) for value in values if normalize_text(value))
    for pattern in (r"(?:受理号码|受理号|号码|手机号|联系电话)\s*[:：]?\s*(1\d{10})", r"(?<!\d)(1\d{10})(?!\d)"):
        match = re.search(pattern, joined)
        if match:
            return match.group(1)
    return None


def extract_employee_name(*values: str) -> str:
    joined = " ".join(normalize_text(value) for value in values if normalize_text(value))
    explicit = re.search(r"(?:员工|员工姓名|专员|客服)\s*[:：]\s*([一-龥·]{2,8})", joined)
    if explicit:
        return normalize_text(explicit.group(1))
    filename_patterns = [
        r"网格\s*([一-龥·]{2,8})\s*[+＋_\-]?\s*1\d{10}",
        r"[年月日号_\-\s]+([一-龥·]{2,8})\s*[+＋_\-]?\s*1\d{10}",
        r"([一-龥·]{2,8})\s*[+＋_\-]?\s*1\d{10}",
    ]
    for pattern in filename_patterns:
        match = re.search(pattern, joined)
        if match:
            candidate = normalize_text(match.group(1))
            if candidate and candidate not in {"网格", "移动", "高清", "客户", "号码"}:
                return candidate
    return ""


def extract_recording_date(*values: str) -> str | None:
    joined = " ".join(normalize_text(value) for value in values if normalize_text(value))
    patterns = [
        r"((?:20)?\d{2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*(?:日|号)?",
        r"((?:20)?\d{2})[-_.](\d{1,2})[-_.](\d{1,2})",
        r"((?:20)?\d{2})(\d{2})(\d{2})",
    ]
    for pattern in patterns:
        match = re.search(pattern, joined)
        if not match:
            continue
        year = int(match.group(1))
        if year < 100:
            year += 2000
        month = int(match.group(2))
        day = int(match.group(3))
        try:
            return datetime(year, month, day).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def parse_submission(text: str, rules: dict[str, Any]) -> AudioSubmission:
    text = repair_multiline_text(text)
    trigger = normalize_text(rules.get("trigger", "#听音检测"))
    aliases = {normalize_text(item) for item in rules.get("trigger_aliases", [trigger, "#听音质检", "#录音质检"]) if normalize_text(item)}
    if not any(item in text for item in aliases):
        raise AudioQualityError(f"消息缺少触发词：{trigger}")
    known = {item for item in ("#沟通记录", "#听音检测", "#听音质检", "#录音质检", "#工单催办") if item in text}
    if known - aliases:
        raise AudioQualityError("一条消息只能执行一个功能，请只保留一个触发词")
    employee_name = extract_employee_name(text)
    order_match = re.search(r"(?:工单号|工单流水号|工单编号)\s*[:：]\s*([^\s，,。；;]+)", text)
    order_id = normalize_text(order_match.group(1)) if order_match else None
    if order_id and not re.fullmatch(str(rules.get("order_id_pattern", r"[A-Za-z0-9_-]{4,64}")), order_id):
        raise AudioQualityError(f"工单号格式不合法：{order_id}")
    return AudioSubmission(employee_name, order_id, extract_accepted_number(text))


def validate_attachment(message: IncomingAudioMessage, rules: dict[str, Any]) -> None:
    if not normalize_text(message.attachment_id):
        raise AudioQualityError("消息缺少音频附件")
    if not normalize_text(message.filename) or Path(message.filename).name != message.filename:
        raise AudioQualityError("音频文件名为空或包含路径")
    suffix = Path(repair_text(message.filename)).suffix.casefold()
    supported = {str(item).casefold() for item in rules.get("supported_audio_suffixes", [])}
    if suffix not in supported:
        raise AudioQualityError(f"不支持的音频格式：{suffix or '<无扩展名>'}")
    if message.file_size is not None:
        maximum = int(rules.get("max_file_size_bytes", 200 * 1024 * 1024))
        if message.file_size <= 0:
            raise AudioQualityError("音频文件为空")
        if message.file_size > maximum:
            raise AudioQualityError("音频文件超过大小限制")


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    if not path.is_file():
        raise AudioQualityError(f"音频文件不存在：{path}")
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(chunk_size), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_transcript(payload: Any, duration_seconds: float | None = None) -> list[TranscriptSegment]:
    if not isinstance(payload, list) or not payload:
        raise AudioQualityError("转写结果必须是非空数组")
    result: list[TranscriptSegment] = []
    previous_start = -1.0
    for index, item in enumerate(payload, start=1):
        if not isinstance(item, dict):
            raise AudioQualityError(f"转写片段 {index} 必须是对象")
        try:
            start = float(item["start_seconds"])
            end = float(item["end_seconds"])
        except (KeyError, TypeError, ValueError) as exc:
            raise AudioQualityError(f"转写片段 {index} 时间格式不合法") from exc
        speaker = normalize_text(item.get("speaker"))
        text = normalize_text(item.get("text"))
        if start < 0 or end <= start or start < previous_start:
            raise AudioQualityError(f"转写片段 {index} 时间范围或顺序不合法")
        if duration_seconds is not None and end > duration_seconds + 1:
            raise AudioQualityError(f"转写片段 {index} 超出音频时长")
        if not speaker or not text:
            raise AudioQualityError(f"转写片段 {index} 缺少说话人或文本")
        result.append(TranscriptSegment(start, end, speaker, text))
        previous_start = start
    return result


def _keyword_score(text: str, keywords: list[str]) -> int:
    return sum(1 for keyword in keywords if keyword and keyword in text)


def evaluate_asr_quality(segments: list[TranscriptSegment], rules: dict[str, Any]) -> list[str]:
    config = rules.get("asr_quality")
    if not isinstance(config, dict) or not config.get("enabled", True):
        return []
    text = normalize_text(" ".join(segment.text for segment in segments))
    reasons: list[str] = []
    min_chars = int(config.get("min_total_characters", 0))
    if min_chars and len(text) < min_chars:
        reasons.append(f"ASR有效文本过短（{len(text)}字，低于{min_chars}字）")
    min_segments = int(config.get("min_segment_count", 0))
    if min_segments and len(segments) < min_segments:
        reasons.append(f"ASR片段数量过少（{len(segments)}段，低于{min_segments}段）")
    unknown = {normalize_text(item) for item in config.get("unknown_speakers", ["", "??", "ASR", "UNKNOWN", "unknown"])}
    speakers = [normalize_text(s.speaker) for s in segments]
    distinct = {speaker for speaker in speakers if speaker not in unknown}
    min_speakers = int(config.get("min_distinct_speakers", 0))
    require_separation = bool(config.get("require_speaker_separation", False))
    if require_separation and min_speakers and len(distinct) < min_speakers:
        reasons.append(f"ASR可用说话人标签少于{min_speakers}个")
    if speakers:
        max_unknown_ratio = float(config.get("max_unknown_speaker_ratio", 1.0))
        if max_unknown_ratio < 1.0 and sum(1 for s in speakers if s in unknown) / len(speakers) > max_unknown_ratio:
            reasons.append("ASR说话人分离不可靠")
    patterns = [normalize_text(item) for item in config.get("uncertain_patterns", []) if normalize_text(item)]
    if any(re.search(pattern, text, flags=re.IGNORECASE) for pattern in patterns):
        reasons.append("ASR文本包含听不清或无法识别标记")
    terms = [normalize_text(item) for item in config.get("key_business_terms", []) if normalize_text(item)]
    if terms and not any(term in text for term in terms):
        reasons.append("ASR未识别到套餐、资费、办理等关键业务信息")
    return reasons


def choose_employee_speaker(segments: list[TranscriptSegment], rules: dict[str, Any]) -> tuple[str, float]:
    if not segments:
        raise AudioQualityError("缺少转写文本，无法判断客服声道")
    opening = [normalize_text(item) for item in rules.get("service_opening_keywords", []) if normalize_text(item)]
    general = [normalize_text(item) for item in rules.get("employee_speaker_keywords", []) if normalize_text(item)]
    opening = opening or ["中国移动", "10086", "工号", "您好", "机主本人", "后四位", "我帮您查询", "请稍等"]
    general = general or ["中国移动", "10086", "工号", "您好", "为您查询", "推荐", "建议", "请稍等"]
    speakers = sorted({segment.speaker for segment in segments})
    if speakers == ["??"]:
        return "??", 0.95
    stats = {speaker: {"opening": 0.0, "early": 0.0, "general": 0.0, "chars": 0.0} for speaker in speakers}
    for segment in segments:
        opening_hits = _keyword_score(segment.text, opening)
        general_hits = _keyword_score(segment.text, general)
        stats[segment.speaker]["chars"] += len(segment.text)
        stats[segment.speaker]["general"] += general_hits
        if segment.start_seconds <= float(rules.get("service_opening_window_seconds", 90)):
            stats[segment.speaker]["early"] += general_hits
            stats[segment.speaker]["opening"] += opening_hits
    ranked = []
    for speaker, item in stats.items():
        ranked.append((speaker, item["opening"] * 12 + item["early"] * 4 + item["general"] * 1.5))
    ranked.sort(key=lambda item: (item[1], stats[item[0]]["opening"], stats[item[0]]["early"], stats[item[0]]["chars"]), reverse=True)
    best, best_score = ranked[0]
    second = ranked[1][1] if len(ranked) > 1 else 0.0
    if stats[best]["opening"] <= 0 and stats[best]["early"] < 2:
        return best, 0.35
    confidence = 0.58 + min(0.25, stats[best]["opening"] * 0.08) + min(0.12, (best_score - second) * 0.03)
    if stats[best]["opening"] <= 0:
        confidence = min(confidence, 0.70)
    return best, min(0.95, confidence)


def scan_rule_hits(segments: Iterable[TranscriptSegment], employee_speaker: str, rules: dict[str, Any]) -> list[RuleHit]:
    hits: list[RuleHit] = []
    seen: set[tuple[str, float, str]] = set()
    for segment in segments:
        if segment.speaker != employee_speaker:
            continue
        for rule in rules.get("phrase_rules", []):
            pattern = normalize_text(rule.get("pattern"))
            if not rule.get("enabled", True) or not pattern:
                continue
            matched = re.search(pattern, segment.text) if rule.get("regex") else pattern in segment.text
            exclude_patterns = [normalize_text(item) for item in rule.get("exclude_patterns", []) if normalize_text(item)]
            excluded = any(re.search(item, segment.text) for item in exclude_patterns)
            key = (str(rule.get("id")), segment.start_seconds, segment.text)
            if matched and not excluded and key not in seen:
                seen.add(key)
                hits.append(RuleHit(str(rule["id"]), segment.start_seconds, segment.text, normalize_text(rule.get("description")), max(0, int(rule.get("deduction", 0)))))
    return hits


SEMANTIC_ACTION_TEMPLATES = {
    "tv_box": {
        "missing_ask": {"scope": "full", "patterns": ["有什么疑问", "想不用", "不用了吗", "因为扣钱", "一直都没有在使用", "还继续用"]},
        "missing_check": {"scope": "full", "patterns": ["合约", "扣", "包含", "到期", "优惠", "15块", "88套餐", "申请"]},
        "missing_compare": {"scope": "full", "patterns": ["优惠", "15块", "一年", "继续使用", "退掉", "88", "78"]},
        "missing_calculate": {"scope": "full", "patterns": ["15块", "一年", "每个月", "88", "78", "优惠"]},
        "missing_retention_action": {"scope": "employee", "patterns": ["优惠", "继续使用", "申请", "帮你申请", "给你免", "减免"]},
    },
    "mobile_plan": {
        "missing_check": {"scope": "full", "patterns": ["套餐", "资费", "月租", "流量", "分钟", "低消", "保号"]},
        "missing_compare": {"scope": "full", "patterns": ["8元", "18元", "最低", "保号", "降到", "改成", "当前套餐"]},
        "missing_calculate": {"scope": "full", "patterns": ["元", "块", "月租", "费用", "每月", "差", "省"]},
        "missing_retention_action": {"scope": "employee", "patterns": ["推荐", "建议", "办理", "保留", "优惠", "申请"]},
    },
}


PROCESS_ACTION_RULE_IDS = {"missing_ask", "missing_check", "missing_compare", "missing_calculate"}

DEFAULT_RETENTION_SUCCESS_PATTERNS = [
    r"(那我先|我先).{0,12}(试一下|用一下|继续用|保留)",
    r"继续保留使用|保留[、,， ]*保留|问题已经解决|已经处理好|已经解决了",
]


def resolve_retention_result(retention_result: str | None, call_summary: str | None = None) -> str:
    summary = normalize_text(call_summary)
    if re.search(r"客户.{0,12}(同意|接受|愿意|决定).{0,12}(优惠|方案|办理|申请|继续使用|保留)", summary):
        return "客户同意接受处理方案"
    return normalize_text(retention_result)

DEFAULT_ACTION_EVIDENCE_PATTERNS = {
    "missing_ask": [
        r"什么.{0,12}(原因|问题|情况|疑问|需求)",
        r"(原因|问题|情况|疑问|需求).{0,12}(是什么|什么|吗|呢)",
        r"是什么原因(呢|吗)?",
        r"(有什么疑问|有啥疑问|有什么问题|有啥问题)",
        r"(投诉|反馈|工单).{0,30}(问题|情况|原因|事项|诉求).{0,12}(吗|呢|是吗|对吗)",
        r"(您|你).{0,24}(投诉|反馈|打过10086|打过一零零八六).{0,30}(问题|情况|原因|事项|诉求).{0,12}(吗|呢|是吗|对吗)",
        r"(看到|收到|见到).{0,24}(投诉|反馈|反映|工单).{0,40}(问题|宽带|机顶盒|移动高清|套餐|业务).{0,16}(吗|呢|是吗|是吧|对吗)",
        r"(您|你).{0,12}(是要|想|需要|打算).{0,20}(取消|不用|退|停|拆).{0,12}(吗|呢|是吗|是吧)",
        r"(地址|小区|位置|在哪里|是哪|哪里).{0,16}(吗|呢|是吗|是吧)?",
        r"(您|你).{0,20}(想|需要|打算|准备|觉得|是因为|主要是|是不是|有没有|是否)",
        r"(为什么|怎么|哪里|哪儿|哪方面|什么情况|有什么疑问|想不用|不用了吗|还继续用吗)",
        r"(经常出差|本地市|外地|家里人|老人|小孩|宽带|信号|流量|费用|套餐).{0,12}(吗|呢|怎么样|好不好)",
    ],
    "missing_check": [
        r"(查|查询|查一下|看一下|核实|确认|后台|系统|显示|看到|帮您看|帮你看)",
        r"(套餐|资费|月租|消费|流量|分钟|语音|合约|低消|宽带|增值业务|机顶盒|移动高清|扣费|优惠|到期|生效|验证码)",
    ],
    "missing_compare": [
        r"(相比|对比|比|原来|之前|现在|当前|新套餐|老套餐|原套餐|现套餐|这个套餐|那个套餐)",
        r"(8元|18元|最低|保号|流量|分钟|月租|费用|优惠|一年|继续使用|退掉|降到|改成)",
    ],
    "missing_calculate": [
        r"(\d{1,4})\s*(元|块|块钱|G|GB|分钟|个月|年)",
        r"(每月|每个月|一个月|一年|合计|总共|折后|差额|优惠|减免|返还|到期|生效|费用|资费|月租)",
    ],
    "missing_retention_action": [
        r"(推荐|建议|可以办理|给您办理|帮您办理|保留|优惠|方案|申请|反馈|继续使用|减免|免掉|考虑一下|工作人员联系)",
    ],
}


def _semantic_action_satisfied(rule_id: str, employee_text: str, full_text: str, semantic_context: list[dict[str, Any]]) -> bool:
    if not rule_id or not semantic_context:
        return False
    for item in semantic_context:
        if not isinstance(item, dict):
            continue
        slots = item.get("slots") if isinstance(item.get("slots"), dict) else {}
        object_id = normalize_text(slots.get("business_object"))
        template = SEMANTIC_ACTION_TEMPLATES.get(object_id, {}).get(rule_id)
        if not template:
            continue
        haystack = employee_text if template.get("scope") == "employee" else full_text
        patterns = [normalize_text(pattern) for pattern in template.get("patterns", []) if normalize_text(pattern)]
        if any(pattern in haystack for pattern in patterns):
            return True
    return False


EVIDENCE_EXCERPT_MAX_CHARS = 160


def _contains_any_pattern(text: str, patterns: Iterable[str]) -> bool:
    return any(re.search(pattern, text, flags=re.IGNORECASE) for pattern in patterns if pattern)


def _evidence_anchor_patterns(rule_id: str) -> list[str]:
    anchors = {
        "missing_ask": [
            r"请问|想问|为什么|什么原因|什么问题|什么情况|有什么|是否|是不是|方便|吗|呢",
            r"投诉|反馈|工单|宽带|套餐|费用|取消|不用|移动高清|机顶盒",
        ],
        "missing_check": [
            r"查询|查一下|查到|找到|核实|系统|后台|看到|显示|验证码|动态码|短信码",
            r"清单|工单|套餐|费用|资费|流量|合约|权限|营业厅|地址|增值业务费|基础包|权益包|生活权益包|话费券|中国移动APP|APP|权益会",
            r"附近|当地|本地|那边|地址|小区|区域|身份证|设备|营业厅|厅",
        ],
        "missing_compare": [
            r"方案|原来|之前|现在|当前|改成|改为|调整|降到|优惠|减免|补退|退费|两个",
            r"\d{1,4}\s*(元|块|块钱|G|GB|分钟)",
        ],
        "missing_calculate": [
            r"每月|每个月|费用|资费|月租|优惠|减免|补退|退费|违约金|生效|到期|一年|两年|增值业务费|基础包|权益包|生活权益包|话费券|中国移动APP|APP|权益会",
            r"\d{1,4}\s*(元|块|块钱|G|GB|分钟|个月|年)",
        ],
    }
    return anchors.get(rule_id, [])


def _evidence_excerpt(text: str, rule_id: str) -> str:
    normalized = normalize_text(text)
    if len(normalized) <= EVIDENCE_EXCERPT_MAX_CHARS:
        return normalized
    anchor_at: int | None = None
    for pattern in _evidence_anchor_patterns(rule_id):
        match = re.search(pattern, normalized, flags=re.IGNORECASE)
        if match:
            anchor_at = match.start()
            break
    if anchor_at is None:
        anchor_at = 0
    start = max(0, anchor_at - 55)
    end = min(len(normalized), anchor_at + 105)
    for delimiter in ("。", "；", ";", "！", "？", "?", "!", "\n"):
        previous = normalized.rfind(delimiter, 0, anchor_at)
        if previous >= 0 and previous + 1 > start:
            start = previous + 1
            break
    for delimiter in ("。", "；", ";", "！", "？", "?", "!", "\n"):
        following = normalized.find(delimiter, anchor_at, len(normalized))
        if following >= 0 and following + 1 < end:
            end = following + 1
            break
    excerpt = normalized[start:end].strip()
    if len(excerpt) > EVIDENCE_EXCERPT_MAX_CHARS:
        excerpt = excerpt[:EVIDENCE_EXCERPT_MAX_CHARS].rstrip()
    return ("..." if start > 0 else "") + excerpt + ("..." if end < len(normalized) else "")


def _has_service_action_cue(rule_id: str, text: str) -> bool:
    service_cues = {
        "missing_ask": [
            r"我.*(想问|请问|了解|核实)|这边.*(想问|了解|核实)",
            r"(您|你).{0,18}(什么原因|什么问题|什么情况|是否|是不是|方便|还需要|继续使用|是要|想取消|要取消|不用|地址|哪里)",
            r"投诉|反馈|反映|工单|业务问题|移动高清|机顶盒|宽带|有什么疑问|是什么原因",
        ],
        "missing_check": [
            r"我|我们|这边|后台|系统|客服|工作人员",
            r"查询|查一下|查到|找到|看到|显示|核实|验证码|动态码|短信码|清单",
            r"(可以去|建议去|您可以去|你可以去|发.*地址|搜.*地址|有权限|权限.*办理|权限.*取消)",
        ],
        "missing_compare": [
            r"我|我们|这边|帮您|帮你|给您|给你|可以|建议|方案",
            r"改成|改为|调整|优惠|减免|补退|退费|回收|申请|上单|继续使用",
        ],
        "missing_calculate": [
            r"我|我们|这边|帮您|帮你|给您|给你|可以|建议|方案",
            r"每月|每个月|费用|优惠|减免|补退|退费|违约金|生效|免费|不用钱|不再扣|增值业务费|基础包|权益包|生活权益包|话费券|权益会",
        ],
        "missing_retention_action": [
            r"我|我们|这边|帮您|帮你|给您|给你|要不|建议",
            r"申请|优惠|减免|免费|0元|不用钱|继续使用|继续用|保留|改成|改为|上单|补退|回收|取消|处理",
        ],
    }
    return _contains_any_pattern(text, service_cues.get(rule_id, []))


def _has_customer_perspective(text: str) -> bool:
    customer_patterns = [
        r"^(我|俺|我们|人家).{0,36}(想|要|不要|不想|不需要|不用|不去了|取消|退掉|回收|投诉|反馈|反映|改|换|降|保号|没用|人在外地|不在|没定|需要改)",
        r"(我|俺|我们).{0,36}(没用|不用|不需要|不接受|不同意|坚持|投诉|反馈|反映|取消|退掉|人在外地|不在|没定|要改|需要改)",
        r"(你们|你这边|移动公司|移动).{0,36}(给我|给我们|帮我|帮我们|怎么|咋|哪里|哪个|去哪|什么时候|可以吗|行吗|退|补退|取消|处理|答复|流程|拖沓|交接|没有结果)",
        r"(我|我们).{0,24}(现在|就是|想问|需要|要).{0,40}(你|你们|移动|客服|工作人员).{0,40}(答复|回答|处理|取消|退|补退)",
        r"你没听懂|你看你回答不出来|明白吗|我才.{0,12}反馈|我又得重新跟你讲",
    ]
    return _contains_any_pattern(normalize_text(text), customer_patterns)


def _is_likely_customer_only_evidence(rule_id: str, text: str) -> bool:
    normalized = normalize_text(text)
    if not normalized:
        return True
    if _has_service_action_cue(rule_id, normalized):
        return False
    customer_only_patterns = [
        r"^(我|俺|我们|人家).{0,24}(想|要|不要|不想|不需要|不用|不去了|取消|退掉|回收|投诉|改|换|降|保号)",
        r"^(那|这个|你|你们).{0,24}(给我|帮我|怎么|咋|哪里|哪个|去哪|什么时候|可以吗|行吗|退|补退|取消)",
        r"(我|俺|我们).{0,20}(没用|不用|不需要|不接受|不同意|坚持|投诉|取消|退掉)",
        r"(那我|我).{0,16}(去|到).{0,12}(哪个|哪里|哪家).{0,8}(营业厅|厅)",
        r"^(嗯|好的|可以|行|知道了|不要了|不用了)[。！？,，;；]*$",
    ]
    return _contains_any_pattern(normalized, customer_only_patterns)


def _evidence_body(text: str) -> str:
    normalized = normalize_text(text)
    normalized = re.sub(r"^已体现[:：]", "", normalized)
    normalized = re.sub(r"^未体现[:：]", "", normalized)
    normalized = re.sub(r"^\[\d+(?:\.\d+)?-\d+(?:\.\d+)?\]\s*[^：:]{1,12}[：:]", "", normalized)
    return normalized.strip()


def _overlaps_customer_segment(text: str, customer_segments: Iterable[TranscriptSegment] | None = None) -> bool:
    body = _evidence_body(text)
    if not body or not customer_segments:
        return False
    for segment in customer_segments:
        customer_text = normalize_text(segment.text)
        if len(customer_text) < 8:
            continue
        if body in customer_text or customer_text in body:
            return True
        # ASR excerpts may trim punctuation and leading fillers; use a conservative overlap.
        limit = min(len(body), len(customer_text))
        if limit >= 18:
            for size in range(min(limit, 42), 17, -1):
                if any(body[start:start + size] in customer_text for start in range(0, max(1, len(body) - size + 1), max(1, size // 2))):
                    return True
    return False


def filter_customer_perspective_evidence(
    evidence: str,
    customer_segments: Iterable[TranscriptSegment] | None = None,
) -> str:
    raw_value = str(evidence or "").strip()
    normalized = normalize_text(raw_value)
    if not normalized or normalized == "未找到明确证据":
        return normalized or "未找到明确证据"
    status_prefix = ""
    raw_body = raw_value
    for prefix in ("已体现：", "未体现：", "已体现:", "未体现:"):
        if raw_body.startswith(prefix):
            status_prefix = prefix
            raw_body = raw_body[len(prefix):]
            break
    separator = "\n" if "\n" in raw_body else "；"
    pieces = [normalize_text(piece) for piece in raw_body.split(separator) if normalize_text(piece)]
    kept = [
        piece
        for piece in pieces
        if not _overlaps_customer_segment(piece, customer_segments)
    ]
    if kept:
        normalized_prefix = status_prefix.replace(":", "：")
        return normalized_prefix + separator.join(kept)
    replacement = "证据疑似客户表述，已在输出中过滤，需人工核验"
    return status_prefix + replacement if status_prefix else "未找到明确证据"


def _format_evidence(segment: TranscriptSegment, rule_id: str = "") -> str:
    text = _evidence_excerpt(segment.text, rule_id) if rule_id else normalize_text(segment.text)
    return f"[{segment.start_seconds:.1f}-{segment.end_seconds:.1f}] {segment.speaker}：{text}"


def _semantic_action_evidence(
    rule_id: str,
    segments: list[TranscriptSegment],
    employee_speaker: str | None,
    rules: dict[str, Any] | None,
    max_items: int,
) -> list[str] | None:
    semantic_config = (rules or {}).get("action_semantics")
    if not isinstance(semantic_config, dict) or not semantic_config.get("enabled", False):
        return None
    if rule_id not in PROCESS_ACTION_RULE_IDS:
        return None

    # Speaker diarization is noisy on many mono recordings; treat speaker labels as a weak hint only.
    usable = list(segments)
    evidence: list[str] = []
    seen: set[str] = set()
    seen_lines: set[str] = set()
    reason_question = re.compile(
        r"(为什么|什么原因|因为什么|什么需求|什么诉求|怎么会|为什么想).{0,24}(改|换|降|取消|不用|退|保留|宽带|套餐|业务)|"
        r"(?:改|换|降|取消|不用|退|保留|宽带|套餐|业务).{0,24}(为什么|什么原因|因为什么|什么需求|什么诉求)|"
        r"是什么原因(呢|吗)?|有什么疑问|有什么问题|"
        r"(投诉|反馈|工单).{0,30}(问题|情况|原因|事项|诉求).{0,12}(吗|呢|是吗|对吗)|"
        r"(您|你).{0,24}(投诉|反馈|打过10086|打过一零零八六).{0,30}(问题|情况|原因|事项|诉求).{0,12}(吗|呢|是吗|对吗)|"
        r"(看到|收到|见到).{0,24}(投诉|反馈|反映|工单).{0,40}(问题|宽带|机顶盒|移动高清|套餐|业务).{0,16}(吗|呢|是吗|是吧|对吗)|"
        r"(您|你).{0,12}(是要|想|需要|打算).{0,20}(取消|不用|退|停|拆).{0,12}(吗|呢|是吗|是吧)|"
        r"(地址|小区|位置|在哪里|是哪|哪里).{0,16}(吗|呢|是吗|是吧)?"
    )
    lookup_action = re.compile(r"查询|查一下|查到|查到了|核实|系统|验证码|动态码|短信码|找到了|后台显示|我这里看到|这边看到")
    account_fact = re.compile(r"(套餐|流量|语音|话费|费用|消费|增值业务|业务费|增值业务费|宽带|合约|低消|权限|月租|权益包|生活权益包|基础包|移动高清|话费券|中国移动APP|APP|权益会|工单|投诉|反馈).{0,18}(\d{1,4}\s*(元|块|G|GB|分钟)|当前|本月|近.{0,3}月|使用|到期|生效|问题|记录|状态|可以|领取|领|查询|搜索)|\d{1,4}\s*(元|块|G|GB|分钟).{0,18}(套餐|流量|语音|话费|费用|业务|增值业务|业务费|宽带|合约|权益|权益包|基础包|话费券)")
    office_lookup = re.compile(
        r"(附近|当地|本地|那边|地址|小区|区域).{0,35}(营业厅|[\u4e00-\u9fa5]{1,8}厅).{0,35}(取消|办理|处理|权限|地址|退|回收|申请)|"
        r"(营业厅|[\u4e00-\u9fa5]{1,8}厅).{0,35}(取消|办理|处理|权限|地址|退|回收|申请|可以|能)|"
        r"(身份证|设备|光猫|路由器).{0,35}(营业厅|[\u4e00-\u9fa5]{1,8}厅|取消|退|回收|交还|办理)|"
        r"(权限).{0,24}(不一定有|有|没有|办理|取消|处理)"
    )
    offer = re.compile(r"(推荐|建议|给.{0,8}优惠|帮.{0,8}(申请|改|办理)|可以.{0,8}(改|办|保留)|改成|改为|调整为|降到|方案)")
    comparison = re.compile(r"(原来|原套餐|现在|当前|之前|改成|改为|调整为|降到|优惠|省|少.{0,4}钱|多.{0,4}(流量|分钟|权益)|免费|0元|两.{0,4}方案|两个方案|\d{1,4}\s*(元|块|G|GB|分钟))")
    benefit = re.compile(r"(优惠.{0,20}\d{1,4}\s*(元|块)|省.{0,12}\d{1,4}\s*(元|块)|免费|不用钱|不收费|不会再扣|免.{0,4}违约金|减免.{0,4}违约金|继续使用|不影响|可领.{0,12}(话费券|权益)|优惠一年|减免|改成.{0,12}\d{1,4}\s*(元|块)|每(月|个?月).{0,12}(加|只要|仅需).{0,8}\d{1,4}\s*(元|块))")
    fee_breakdown = re.compile(r"(增值业务费|业务费|套餐费|月租|基础包|权益包|生活权益包|移动高清|话费券).{0,18}(\d{1,4}\s*(元|块|块钱)|可以领|领取|权益|费用)|\d{1,4}\s*(元|块|块钱).{0,18}(增值业务|业务费|套餐|月租|基础包|权益包|生活权益包|话费券)|(中国移动APP|APP|权益会).{0,28}(查询|搜索|登录|领|领取|话费券|权益)")

    for index, segment in enumerate(usable):
        text = normalize_text(segment.text)
        if not text:
            continue
        if _is_likely_customer_only_evidence(rule_id, text):
            continue
        start = max(0, index - 4)
        end = min(len(usable), index + 3)
        context_segments = usable[start:end]
        context = " ".join(normalize_text(item.text) for item in context_segments)
        matched = False
        matched_segments = [segment]
        if rule_id == "missing_ask":
            matched = bool(reason_question.search(text))
        elif rule_id == "missing_check":
            # A check requires both an observable lookup and an account-specific result.
            matched = bool((account_fact.search(text) and lookup_action.search(context)) or office_lookup.search(text))
            if matched:
                matched_segments = [
                    item
                    for item in context_segments
                    if not _is_likely_customer_only_evidence(rule_id, item.text)
                    and (
                        lookup_action.search(normalize_text(item.text))
                        or account_fact.search(normalize_text(item.text))
                        or office_lookup.search(normalize_text(item.text))
                    )
                ] or [segment]
        elif rule_id == "missing_compare":
            matched = bool(offer.search(text) and comparison.search(text))
        elif rule_id == "missing_calculate":
            matched = bool((offer.search(text) and benefit.search(text)) or fee_breakdown.search(text))
        if not matched:
            continue
        formatted_lines = [_format_evidence(item, rule_id) for item in matched_segments]
        new_lines = [line for line in formatted_lines if line not in seen_lines]
        if not new_lines:
            continue
        formatted = "\n".join(new_lines)
        if formatted not in seen:
            seen.add(formatted)
            seen_lines.update(new_lines)
            evidence.append(formatted)
            if len(evidence) >= max_items:
                break
    return evidence


def _find_action_evidence(
    rule_id: str,
    segments: Iterable[TranscriptSegment],
    employee_speaker: str | None = None,
    rules: dict[str, Any] | None = None,
    max_items: int = 3,
) -> list[str]:
    segment_list = list(segments)
    semantic_evidence = _semantic_action_evidence(rule_id, segment_list, employee_speaker, rules, max_items)
    if semantic_evidence:
        return semantic_evidence
    configured = (rules or {}).get("action_evidence_rules")
    configured_patterns: list[str] = []
    if isinstance(configured, dict):
        item = configured.get(rule_id)
        if isinstance(item, dict):
            configured_patterns = [
                normalize_text(pattern)
                for pattern in item.get("patterns", [])
                if normalize_text(pattern)
            ]
    patterns = configured_patterns or DEFAULT_ACTION_EVIDENCE_PATTERNS.get(rule_id, [])
    if not patterns:
        return []
    evidence: list[str] = []
    seen: set[str] = set()
    for index, segment in enumerate(segment_list):
        text = normalize_text(segment.text)
        if not text:
            continue
        if _is_likely_customer_only_evidence(rule_id, text):
            continue
        matched_segments = [segment]
        matched_text = text
        if rule_id == "missing_check":
            start = max(0, index - 1)
            end = min(len(segment_list), index + 2)
            window = [
                item for item in segment_list[start:end]
                if not _is_likely_customer_only_evidence(rule_id, item.text)
            ]
            window_text = " ".join(normalize_text(item.text) for item in window)
            if any(re.search(pattern, window_text, flags=re.IGNORECASE) for pattern in patterns):
                matched_segments = window
                matched_text = window_text
        if any(re.search(pattern, matched_text, flags=re.IGNORECASE) for pattern in patterns):
            formatted = "\n".join(_format_evidence(item, rule_id) for item in matched_segments)
            if formatted not in seen:
                seen.add(formatted)
                evidence.append(formatted)
                if len(evidence) >= max_items:
                    break
    return evidence


def find_process_action_evidence(
    segments: Iterable[TranscriptSegment],
    employee_speaker: str | None = None,
    all_segments: Iterable[TranscriptSegment] | None = None,
    rules: dict[str, Any] | None = None,
) -> dict[str, str]:
    segment_list = list(segments)
    all_segment_list = list(all_segments) if all_segments is not None else segment_list
    result: dict[str, str] = {}
    for rule_id in ["missing_ask", "missing_check", "missing_compare", "missing_calculate"]:
        primary = _find_action_evidence(rule_id, segment_list, employee_speaker, rules)
        fallback = [] if primary else _find_action_evidence(rule_id, all_segment_list, None, rules)
        result[rule_id] = "\n".join(primary or fallback) or "未找到明确证据"
    return result


def scan_required_action_gaps(
    segments: Iterable[TranscriptSegment],
    employee_speaker: str,
    rules: dict[str, Any],
    semantic_context: list[dict[str, Any]] | None = None,
    all_segments: Iterable[TranscriptSegment] | None = None,
    retention_result: str | None = None,
) -> list[RuleHit]:
    segment_list = list(segments)
    all_segment_list = list(all_segments) if all_segments is not None else segment_list
    employee_text = "\n".join(segment.text for segment in segment_list if segment.speaker == employee_speaker)
    full_text = "\n".join(segment.text for segment in all_segment_list)
    semantic_items = semantic_context or []
    evidence_only = bool((rules.get("required_action_scoring") or {}).get("evidence_only", False))
    gaps: list[RuleHit] = []
    for action in rules.get("required_actions", []):
        rule_id = normalize_text(action.get("id"))
        patterns = [normalize_text(item) for item in action.get("patterns", []) if normalize_text(item)]
        if evidence_only and rule_id in PROCESS_ACTION_RULE_IDS:
            continue
        if rule_id == "missing_retention_success":
            success_text = normalize_text(retention_result)
            success_patterns = [normalize_text(item) for item in action.get("success_patterns", []) if normalize_text(item)]
            failure_patterns = [normalize_text(item) for item in action.get("failure_patterns", []) if normalize_text(item)]
            transcript_success_patterns = [
                normalize_text(item)
                for item in action.get("transcript_success_patterns", DEFAULT_RETENTION_SUCCESS_PATTERNS)
                if normalize_text(item)
            ]
            if any(re.search(pattern, full_text) for pattern in transcript_success_patterns):
                continue
            if success_text and failure_patterns and any(re.search(pattern, success_text) for pattern in failure_patterns):
                gaps.append(RuleHit(rule_id, 0.0, success_text, normalize_text(action.get("description")) or normalize_text(action.get("label")), max(0, int(action.get("deduction", 0)))))
                continue
            if success_text and success_patterns and any(re.search(pattern, success_text) for pattern in success_patterns):
                continue
        configured_evidence = rules.get("action_evidence_rules") if isinstance(rules.get("action_evidence_rules"), dict) else {}
        if isinstance(configured_evidence, dict) and rule_id in configured_evidence:
            # Process evidence is evaluated from the complete transcript because role labels can be unreliable on mono audio.
            evidence = _find_action_evidence(rule_id, all_segment_list, None, rules)
            if evidence:
                continue
            gaps.append(RuleHit(rule_id, 0.0, "未找到明确证据", normalize_text(action.get("description")) or normalize_text(action.get("label")), max(0, int(action.get("deduction", 0)))))
            continue
        if rule_id and _semantic_action_satisfied(rule_id, employee_text, full_text, semantic_items):
            continue
        if rule_id and patterns and not any(re.search(pattern, employee_text) for pattern in patterns):
            gaps.append(RuleHit(rule_id, 0.0, "未检测到明确话术", normalize_text(action.get("description")) or normalize_text(action.get("label")), max(0, int(action.get("deduction", 0)))))
    return gaps


def deterministic_quality_check(
    segments: list[TranscriptSegment],
    rules: dict[str, Any],
    employee_speaker_override: str | None = None,
    speaker_confidence_override: float | None = None,
    extra_review_reasons: Iterable[str] | None = None,
    asr_quality_segments: list[TranscriptSegment] | None = None,
    semantic_context: list[dict[str, Any]] | None = None,
    retention_result: str | None = None,
) -> QualityResult:
    if employee_speaker_override:
        employee_speaker = normalize_text(employee_speaker_override)
        if employee_speaker not in {segment.speaker for segment in segments}:
            raise AudioQualityError("指定的员工说话人不存在于待评分片段")
        confidence = float(speaker_confidence_override if speaker_confidence_override is not None else 0.85)
    else:
        employee_speaker, confidence = choose_employee_speaker(segments, rules)
    threshold = float(rules.get("speaker_review_threshold", 0.75))
    hits = [] if confidence < threshold else scan_rule_hits(segments, employee_speaker, rules) + scan_required_action_gaps(segments, employee_speaker, rules, semantic_context, asr_quality_segments or segments, retention_result)
    dimensions = rules.get("dimensions") or {"礼貌规范": {"weight": 0.45, "base_score": 100}, "沟通完整": {"weight": 0.35, "base_score": 90}, "表达清晰": {"weight": 0.20, "base_score": 90}}
    score_policy = rules.get("score_policy") if isinstance(rules.get("score_policy"), dict) else {}
    if str(score_policy.get("mode", "")).lower() in {"uniform", "flat", "flat_deduction"}:
        base_score = int(score_policy.get("base_score", 100))
        score = max(0, min(100, base_score - sum(hit.deduction for hit in hits)))
        dimension_name = str(score_policy.get("dimension_name", "综合质检"))
        scores: dict[str, int] = {dimension_name: score}
    else:
        scores = {}
        for key, config in dimensions.items():
            score = int(config.get("base_score", 100))
            for hit in hits:
                applies = config.get("rule_ids")
                if not applies or hit.rule_id in applies:
                    score -= hit.deduction
            scores[key] = max(0, min(100, score))
        total_weight = sum(max(0.0, float(config.get("weight", 0))) for config in dimensions.values()) or 1.0
        score = int(round(sum(scores[key] * max(0.0, float(config.get("weight", 0))) for key, config in dimensions.items()) / total_weight))
    issues = [QualityIssue(hit.rule_id, hit.start_seconds, hit.quote, hit.description or "命中质检规则", "请按规范话术客观、礼貌沟通，避免引发客户误解。", 1.0) for hit in hits]
    review_reasons = evaluate_asr_quality(asr_quality_segments or segments, rules)
    review_reasons.extend(normalize_text(item) for item in (extra_review_reasons or []) if normalize_text(item))
    needs_review = confidence < threshold or bool(review_reasons)
    if confidence < threshold:
        analysis = "未能通过开头固定话术可靠识别客服声道，未执行自动扣分。"
        optimization = "请人工确认客服声道；确认后再按问、查、算、比、挽留动作进行质检。"
    elif issues:
        analysis = "；".join(f"{_rule_label(issue.rule_id, rules)}：{issue.reason}" for issue in issues)
        optimization = "请根据问题证据复核录音，并按规则库调整沟通话术。"
    else:
        analysis = "未命中已配置的确定性扣分规则。"
        optimization = "保持礼貌、客观、清晰表达；后续可按正式评分标准继续补充规则库。"
    if review_reasons:
        analysis = f"{analysis} 需人工复核：{'；'.join(dict.fromkeys(review_reasons))}。"
        optimization = "请先复核 ASR 转写、客服角色和关键资费/套餐信息；确认后再采纳自动评分。"
    summary = f"员工声道：{employee_speaker}；规则评分：{score}；转写摘要：{join_transcript(segments, max_chars=180)}"
    return QualityResult(employee_speaker, confidence, scores, score, score >= int(rules.get("qualified_score", 80)) and not needs_review, summary, analysis, optimization, issues, needs_review)


def init_db(path: Path) -> sqlite3.Connection:
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path, timeout=10)
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("PRAGMA busy_timeout = 10000")
    connection.execute("PRAGMA journal_mode = WAL")
    connection.executescript("""
        CREATE TABLE IF NOT EXISTS audio_tasks (
            task_id TEXT PRIMARY KEY, message_id TEXT NOT NULL, group_id TEXT NOT NULL,
            sender_user_id TEXT NOT NULL, sender_name TEXT NOT NULL, employee_name TEXT NOT NULL,
            order_id TEXT, accepted_number TEXT, attachment_id TEXT NOT NULL, attachment_hash TEXT NOT NULL,
            original_filename TEXT NOT NULL, file_size INTEGER, duration_seconds REAL,
            status TEXT NOT NULL, created_at TEXT NOT NULL, completed_at TEXT,
            error_stage TEXT, error_message TEXT, attempt_count INTEGER NOT NULL DEFAULT 0,
            claimed_at TEXT, worker_id TEXT, archived_path TEXT, asr_provider TEXT, provider_task_id TEXT,
            UNIQUE (group_id, message_id, attachment_hash));
        CREATE TABLE IF NOT EXISTS audio_segments (
            task_id TEXT NOT NULL, segment_number INTEGER NOT NULL, start_seconds REAL NOT NULL,
            end_seconds REAL NOT NULL, speaker TEXT NOT NULL, text TEXT NOT NULL,
            PRIMARY KEY (task_id, segment_number));
        CREATE TABLE IF NOT EXISTS quality_results (
            task_id TEXT PRIMARY KEY, employee_speaker TEXT NOT NULL, speaker_confidence REAL NOT NULL,
            dimension_scores TEXT NOT NULL, score INTEGER NOT NULL, qualified INTEGER NOT NULL,
            summary TEXT NOT NULL, analysis TEXT NOT NULL, optimization TEXT NOT NULL,
            issues TEXT NOT NULL, needs_human_review INTEGER NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS audio_events (
            event_id INTEGER PRIMARY KEY AUTOINCREMENT, task_id TEXT NOT NULL,
            event_type TEXT NOT NULL, event_data TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS organized_calls (
            task_id TEXT PRIMARY KEY, role_confidence REAL NOT NULL, call_summary TEXT NOT NULL,
            retention_result TEXT NOT NULL, uncertain_parts TEXT NOT NULL, full_dialogue TEXT NOT NULL,
            service_segments TEXT NOT NULL, customer_segments TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS audio_maintenance (
            key TEXT PRIMARY KEY, value TEXT NOT NULL, updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS deepseek_quality_assessments (
            task_id TEXT PRIMARY KEY, score INTEGER NOT NULL, qualified INTEGER NOT NULL,
            summary TEXT NOT NULL, analysis TEXT NOT NULL, issues TEXT NOT NULL,
            actions TEXT NOT NULL, uncertain_parts TEXT NOT NULL, raw_response TEXT NOT NULL,
            created_at TEXT NOT NULL);
    """)
    connection.commit()
    return connection


def clear_audio_tasks(connection: sqlite3.Connection) -> dict[str, int]:
    counts = {
        "audio_tasks": int(connection.execute("SELECT count(*) FROM audio_tasks").fetchone()[0]),
        "audio_segments": int(connection.execute("SELECT count(*) FROM audio_segments").fetchone()[0]),
        "organized_calls": int(connection.execute("SELECT count(*) FROM organized_calls").fetchone()[0]),
        "quality_results": int(connection.execute("SELECT count(*) FROM quality_results").fetchone()[0]),
        "audio_events": int(connection.execute("SELECT count(*) FROM audio_events").fetchone()[0]),
    }
    connection.execute("BEGIN IMMEDIATE")
    connection.execute("DELETE FROM audio_segments")
    connection.execute("DELETE FROM organized_calls")
    connection.execute("DELETE FROM quality_results")
    connection.execute("DELETE FROM audio_events")
    connection.execute("DELETE FROM audio_tasks")
    connection.commit()
    return counts


def run_daily_task_cleanup(connection: sqlite3.Connection, cleanup_config: dict[str, Any] | None, today: str | None = None) -> bool:
    config = cleanup_config or {}
    if not config.get("enabled", False):
        return False
    day = today or datetime.now().strftime("%Y-%m-%d")
    key = "audio_tasks_last_daily_clear"
    row = connection.execute("SELECT value FROM audio_maintenance WHERE key=?", (key,)).fetchone()
    if row and row[0] == day:
        return False
    if config.get("clear_all", False):
        clear_audio_tasks(connection)
    else:
        retention_days = int(config.get("retention_days", 0) or 0)
        if retention_days > 0:
            cutoff = (datetime.now(timezone.utc) - timedelta(days=retention_days)).isoformat(timespec="seconds")
            task_rows = connection.execute(
                """SELECT task_id FROM audio_tasks
                   WHERE status IN ('completed','needs_review','failed') AND created_at < ?""",
                (cutoff,),
            ).fetchall()
            task_ids = [str(row[0]) for row in task_rows]
            if task_ids:
                placeholders = ",".join("?" for _ in task_ids)
                connection.execute("BEGIN IMMEDIATE")
                connection.execute(f"DELETE FROM audio_segments WHERE task_id IN ({placeholders})", task_ids)
                connection.execute(f"DELETE FROM organized_calls WHERE task_id IN ({placeholders})", task_ids)
                connection.execute(f"DELETE FROM quality_results WHERE task_id IN ({placeholders})", task_ids)
                connection.execute(f"DELETE FROM audio_events WHERE task_id IN ({placeholders})", task_ids)
                connection.execute(f"DELETE FROM audio_tasks WHERE task_id IN ({placeholders})", task_ids)
                connection.commit()
    connection.execute(
        """INSERT INTO audio_maintenance(key,value,updated_at) VALUES(?,?,?)
           ON CONFLICT(key) DO UPDATE SET value=excluded.value, updated_at=excluded.updated_at""",
        (key, day, utc_now()),
    )
    connection.commit()
    return True


def add_event(connection: sqlite3.Connection, task_id: str, event_type: str, event_data: dict[str, Any] | None = None) -> None:
    connection.execute("INSERT INTO audio_events(task_id,event_type,event_data,created_at) VALUES(?,?,?,?)", (task_id, event_type, json.dumps(event_data or {}, ensure_ascii=False), utc_now()))


def claim_next_task(connection: sqlite3.Connection, worker_id: str, stale_after_seconds: int = 1800) -> str | None:
    now = datetime.now(timezone.utc)
    stale_before = datetime.fromtimestamp(now.timestamp() - max(1, stale_after_seconds), timezone.utc).isoformat(timespec="seconds")
    connection.execute("BEGIN IMMEDIATE")
    row = connection.execute("""SELECT task_id,status FROM audio_tasks
        WHERE status='received' OR (status='processing' AND (claimed_at IS NULL OR claimed_at < ?))
        ORDER BY created_at,rowid LIMIT 1""", (stale_before,)).fetchone()
    if not row:
        connection.rollback()
        return None
    task_id, previous = str(row[0]), str(row[1])
    connection.execute("""UPDATE audio_tasks SET status='processing',claimed_at=?,worker_id=?,
        attempt_count=attempt_count+1,error_stage=NULL,error_message=NULL WHERE task_id=?""", (utc_now(), worker_id, task_id))
    add_event(connection, task_id, "claimed", {"worker_id": worker_id, "reclaimed": previous == "processing"})
    connection.commit()
    return task_id


def reserve_task(connection: sqlite3.Connection, message: IncomingAudioMessage, submission: AudioSubmission) -> tuple[str, bool]:
    digest = normalize_text(message.attachment_hash) or hashlib.sha256(f"{message.attachment_id}\0{message.filename}\0{message.file_size}".encode()).hexdigest()
    accepted_number = submission.accepted_number or extract_accepted_number(repair_text(message.filename), repair_multiline_text(message.text))
    connection.execute("BEGIN IMMEDIATE")
    existing_by_hash = connection.execute("SELECT task_id FROM audio_tasks WHERE group_id=? AND attachment_hash=? ORDER BY created_at LIMIT 1", (message.group_id, digest)).fetchone()
    if existing_by_hash:
        task_id = str(existing_by_hash[0])
        add_event(connection, task_id, "duplicate_audio_ignored", {"message_id": message.message_id, "filename": message.filename})
        connection.commit()
        return task_id, False
    existing = connection.execute("SELECT task_id FROM audio_tasks WHERE group_id=? AND message_id=? AND attachment_hash=?", (message.group_id, message.message_id, digest)).fetchone()
    if existing:
        connection.rollback()
        return str(existing[0]), False
    task_id = "QA-" + datetime.now().strftime("%Y%m%d") + "-" + uuid.uuid4().hex[:8].upper()
    now = utc_now()
    connection.execute("""INSERT INTO audio_tasks (
        task_id,message_id,group_id,sender_user_id,sender_name,employee_name,order_id,accepted_number,
        attachment_id,attachment_hash,original_filename,file_size,duration_seconds,status,created_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?, 'received',?)""", (task_id, message.message_id, message.group_id, message.sender_user_id, message.sender_name, submission.employee_name, submission.order_id, accepted_number, message.attachment_id, digest, message.filename, message.file_size, message.duration_seconds, now))
    add_event(connection, task_id, "received", asdict(submission))
    connection.commit()
    return task_id, True


def attach_archived_file(connection: sqlite3.Connection, task_id: str, archived_path: Path) -> None:
    resolved = archived_path.resolve()
    if not resolved.is_file():
        raise AudioQualityError(f"归档音频不存在：{resolved}")
    connection.execute("BEGIN IMMEDIATE")
    connection.execute("UPDATE audio_tasks SET archived_path=? WHERE task_id=?", (str(resolved), task_id))
    add_event(connection, task_id, "attachment_archived", {"path": str(resolved)})
    connection.commit()


def mark_task_failed(connection: sqlite3.Connection, task_id: str, stage: str, error: Exception | str) -> None:
    connection.execute("BEGIN IMMEDIATE")
    connection.execute("""UPDATE audio_tasks SET status='failed',error_stage=?,error_message=?,
        completed_at=?,claimed_at=NULL,worker_id=NULL WHERE task_id=?""", (normalize_text(stage), normalize_text(error)[:2000], utc_now(), task_id))
    add_event(connection, task_id, "failed", {"stage": normalize_text(stage), "error": normalize_text(error)[:2000]})
    connection.commit()


def save_transcript(connection: sqlite3.Connection, task_id: str, segments: list[TranscriptSegment]) -> None:
    connection.execute("BEGIN IMMEDIATE")
    connection.execute("DELETE FROM audio_segments WHERE task_id=?", (task_id,))
    connection.executemany("INSERT INTO audio_segments VALUES(?,?,?,?,?,?)", [(task_id, i, s.start_seconds, s.end_seconds, s.speaker, s.text) for i, s in enumerate(segments, 1)])
    connection.execute("UPDATE audio_tasks SET status='transcribed' WHERE task_id=?", (task_id,))
    add_event(connection, task_id, "transcribed", {"segment_count": len(segments)})
    connection.commit()


def save_organized_call(connection: sqlite3.Connection, task_id: str, organized: Any) -> None:
    now = utc_now()
    connection.execute("BEGIN IMMEDIATE")
    connection.execute("""INSERT INTO organized_calls VALUES(?,?,?,?,?,?,?,?,?)
        ON CONFLICT(task_id) DO UPDATE SET role_confidence=excluded.role_confidence,
        call_summary=excluded.call_summary,retention_result=excluded.retention_result,
        uncertain_parts=excluded.uncertain_parts,full_dialogue=excluded.full_dialogue,
        service_segments=excluded.service_segments,customer_segments=excluded.customer_segments,created_at=excluded.created_at""", (task_id, float(organized.role_confidence), normalize_text(organized.call_summary), normalize_text(organized.retention_result), normalize_text(organized.uncertain_parts), str(organized.full_dialogue), json.dumps([asdict(item) for item in organized.service_segments], ensure_ascii=False), json.dumps([asdict(item) for item in organized.customer_segments], ensure_ascii=False), now))
    add_event(connection, task_id, "organized", {"role_confidence": float(organized.role_confidence)})
    connection.commit()


def save_quality_result(connection: sqlite3.Connection, task_id: str, result: QualityResult) -> None:
    now = utc_now()
    status = "needs_review" if result.needs_human_review else "completed"
    connection.execute("BEGIN IMMEDIATE")
    connection.execute("""INSERT INTO quality_results VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
        ON CONFLICT(task_id) DO UPDATE SET employee_speaker=excluded.employee_speaker,
        speaker_confidence=excluded.speaker_confidence,dimension_scores=excluded.dimension_scores,
        score=excluded.score,qualified=excluded.qualified,summary=excluded.summary,
        analysis=excluded.analysis,optimization=excluded.optimization,issues=excluded.issues,
        needs_human_review=excluded.needs_human_review,created_at=excluded.created_at""", (task_id, result.employee_speaker, result.speaker_confidence, json.dumps(result.dimension_scores, ensure_ascii=False), result.score, int(result.qualified), result.summary, result.analysis, result.optimization, json.dumps([asdict(i) for i in result.issues], ensure_ascii=False), int(result.needs_human_review), now))
    connection.execute("UPDATE audio_tasks SET status=?,completed_at=?,claimed_at=NULL,worker_id=NULL WHERE task_id=?", (status, now, task_id))
    add_event(connection, task_id, status, {"score": result.score})
    connection.commit()


def save_deepseek_quality_assessment(connection: sqlite3.Connection, task_id: str, result: QualityResult, assessment: Any) -> None:
    connection.execute(
        """INSERT INTO deepseek_quality_assessments
           (task_id,score,qualified,summary,analysis,issues,actions,uncertain_parts,raw_response,created_at)
           VALUES(?,?,?,?,?,?,?,?,?,?)
           ON CONFLICT(task_id) DO UPDATE SET score=excluded.score,qualified=excluded.qualified,
           summary=excluded.summary,analysis=excluded.analysis,issues=excluded.issues,
           actions=excluded.actions,uncertain_parts=excluded.uncertain_parts,
           raw_response=excluded.raw_response,created_at=excluded.created_at""",
        (
            task_id,
            result.score,
            int(result.qualified),
            result.summary,
            result.analysis,
            json.dumps([asdict(issue) for issue in result.issues], ensure_ascii=False),
            json.dumps(getattr(assessment, "actions", {}), ensure_ascii=False),
            normalize_text(getattr(assessment, "uncertain_parts", "")),
            normalize_text(getattr(assessment, "raw_response", "")),
            utc_now(),
        ),
    )


def load_segments(connection: sqlite3.Connection, task_id: str) -> list[TranscriptSegment]:
    rows = connection.execute("SELECT start_seconds,end_seconds,speaker,text FROM audio_segments WHERE task_id=? ORDER BY segment_number", (task_id,)).fetchall()
    return [TranscriptSegment(float(row[0]), float(row[1]), str(row[2]), str(row[3])) for row in rows]


def join_transcript(segments: Iterable[TranscriptSegment], max_chars: int | None = None) -> str:
    text = "\n".join(f"[{s.start_seconds:.1f}-{s.end_seconds:.1f}] {s.speaker}: {s.text}" for s in segments)
    if max_chars is not None and len(text) > max_chars:
        return text[:max(0, max_chars - 3)] + "..."
    return text


def _issue_rule_ids(result_issues: str | None) -> set[str]:
    if not result_issues:
        return set()
    try:
        raw = json.loads(result_issues)
    except json.JSONDecodeError:
        return set()
    return {normalize_text(item.get("rule_id")) for item in raw if isinstance(item, dict)}


def _rule_label(rule_id: str, rules: dict[str, Any]) -> str:
    normalized = normalize_text(rule_id)
    for action in rules.get("required_actions", []):
        if isinstance(action, dict) and normalize_text(action.get("id")) == normalized:
            return normalize_text(action.get("label")) or normalized
    for rule in rules.get("phrase_rules", []):
        if isinstance(rule, dict) and normalize_text(rule.get("id")) == normalized:
            return normalize_text(rule.get("label")) or normalize_text(rule.get("description")) or normalized
    defaults = {
        "missing_ask": "问",
        "missing_check": "查",
        "missing_compare": "比",
        "missing_calculate": "算",
        "missing_retention_action": "挽留动作",
        "missing_retention_success": "挽留结果",
        "rude_phrase": "服务态度",
        "commanding_tone": "命令式语气",
        "impatient_phrase": "耐心不足",
        "forbidden_no_authority": "推诿风险",
        "push_to_hall_without_retention": "引导营业厅风险",
        "promise_without_basis": "无依据承诺",
        "vague_expression": "表达不清",
        "unclear_result": "处理结果不清",
    }
    return defaults.get(normalized, normalized or "未命名规则")


def _action_status(rule_ids: set[str], missing_rule_id: str) -> str:
    return "未体现" if missing_rule_id in rule_ids else "已体现"


def _weighted_action_status(rule_ids: set[str], missing_rule_id: str) -> str:
    return "0%（未体现）" if missing_rule_id in rule_ids else "25%（已体现）"


def _evidence_weighted_action_status(evidence: str) -> str:
    return "25%（待人工核验）" if not evidence or evidence == "未找到明确证据" else "25%（有证据）"


def _load_deepseek_actions(connection: sqlite3.Connection, task_id: str) -> dict[str, Any] | None:
    row = connection.execute("SELECT actions FROM deepseek_quality_assessments WHERE task_id=?", (task_id,)).fetchone()
    if not row or not row[0]:
        return None
    try:
        actions = json.loads(str(row[0]))
    except json.JSONDecodeError:
        return None
    return actions if isinstance(actions, dict) else None


def _deepseek_action_text(actions: dict[str, Any], rule_id: str) -> str:
    item = actions.get(rule_id)
    if not isinstance(item, dict):
        return "未体现：缺少语义评分结果"
    evidence = [normalize_text(value) for value in item.get("evidence", []) if normalize_text(value)]
    if item.get("present"):
        return "已体现：" + ("；".join(evidence) if evidence else "DeepSeek已判定体现，缺少可展示证据")
    return "未体现：" + (normalize_text(item.get("reason")) or "未找到明确证据")


def _compact_issues(issues: str | None, rules: dict[str, Any] | None = None, exclude_rule_ids: set[str] | None = None) -> str:
    if not issues:
        return ""
    try:
        raw = json.loads(issues)
    except json.JSONDecodeError:
        return normalize_text(issues)
    parts = []
    excluded = exclude_rule_ids or set()
    for item in raw if isinstance(raw, list) else []:
        rule_id = normalize_text(item.get("rule_id")) if isinstance(item, dict) else ""
        if rule_id in excluded:
            continue
        if isinstance(item, dict) and normalize_text(item.get("reason")):
            quote = normalize_text(item.get("quote"))
            label = _rule_label(rule_id, rules or {})
            parts.append(f"{label}：{normalize_text(item.get('reason'))}" + (f"（{quote}）" if quote and quote not in {"未检测到明确话术", "未找到明确证据"} else ""))
    return "；".join(parts)


def _retention_scene_problem(issues: str | None) -> str:
    if not issues:
        return ""
    try:
        raw = json.loads(issues)
    except json.JSONDecodeError:
        return normalize_text(issues)
    scene_rules = {"missing_ask", "missing_check", "missing_calculate", "missing_compare", "missing_retention_action", "missing_retention_success"}
    parts = []
    for item in raw if isinstance(raw, list) else []:
        if isinstance(item, dict) and normalize_text(item.get("rule_id")) in scene_rules:
            parts.append(normalize_text(item.get("reason")) or normalize_text(item.get("rule_id")))
    return "；".join(parts)


def _segments_json_text(value: str | None, role: str) -> str:
    if not value:
        return ""
    try:
        raw = json.loads(value)
    except json.JSONDecodeError:
        return normalize_text(value)
    return "\n".join(f"{role}：{normalize_text(item.get('text'))}" for item in raw if isinstance(item, dict) and normalize_text(item.get("text")))


def _segments_from_json(value: str | None, role: str) -> list[TranscriptSegment]:
    if not value:
        return []
    try:
        raw = json.loads(value)
    except json.JSONDecodeError:
        return []
    segments: list[TranscriptSegment] = []
    for item in raw if isinstance(raw, list) else []:
        if not isinstance(item, dict):
            continue
        text = normalize_text(item.get("text"))
        if not text:
            continue
        try:
            start = float(item.get("start_seconds", 0))
            end = float(item.get("end_seconds", start + 0.001))
        except (TypeError, ValueError):
            start, end = 0.0, 0.001
        if end <= start:
            end = start + 0.001
        segments.append(TranscriptSegment(start, end, role, text))
    return segments


def _infer_service_reason(segments: Iterable[TranscriptSegment], employee_speaker: str | None = None) -> str:
    customer_segments = [segment for segment in segments if not employee_speaker or segment.speaker != employee_speaker]
    text = normalize_text(" ".join(segment.text for segment in customer_segments))
    if not text:
        text = normalize_text(" ".join(segment.text for segment in segments))
    if not text:
        return ""
    candidates = [
        (["太贵", "费用高", "月租高", "套餐贵", "消费高"], "客户认为当前套餐/费用偏高"),
        (["改8元", "8元", "保号", "最低消费", "最低套餐", "降到最低", "降套餐", "降资费", "改低"], "客户希望降低资费或改低档/保号套餐"),
        (["不用流量", "用不上流量", "流量用不完", "老人机", "只接电话"], "客户用量较少或号码主要保留接打电话"),
        (["机顶盒", "移动高清", "魔百和", "不用了", "退掉", "取消", "扣费"], "客户反馈机顶盒/移动高清不用或扣费问题"),
        (["宽带", "网速", "不好用", "卡", "断网"], "客户反馈宽带或网络使用体验问题"),
        (["信号", "没信号", "信号不好", "打不通"], "客户反馈手机信号问题"),
        (["销户", "不用这个号码", "停机", "注销"], "客户有销户/停用号码意向"),
    ]
    reasons = [reason for keywords, reason in candidates if any(keyword in text for keyword in keywords)]
    if reasons:
        return "；".join(dict.fromkeys(reasons))
    question_patterns = [
        r"(因为.{0,30})",
        r"(主要是.{0,30})",
        r"(就是.{0,30})",
        r"(想.{0,30}(套餐|资费|费用|退掉|取消|不用|保号))",
    ]
    for pattern in question_patterns:
        match = re.search(pattern, text)
        if match:
            return normalize_text(match.group(1))
    return text[:120]


def _safe_report_filename(value: str) -> str:
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", normalize_text(value) or "transcript")[:120]


def _write_report_text_file(report_path: Path, task_id: str, suffix: str, text: str) -> Path | None:
    if not text:
        return None
    day = report_path.stem[:10] if re.match(r"\d{4}-\d{2}-\d{2}", report_path.stem) else datetime.now().strftime("%Y-%m-%d")
    folder = report_path.parent / "transcript_texts" / day
    folder.mkdir(parents=True, exist_ok=True)
    output = folder / f"{_safe_report_filename(task_id)}_{suffix}.txt"
    output.write_text(text, encoding="utf-8")
    return output


def export_report(connection: sqlite3.Connection, output: Path, report_day: str | None = None, rules: dict[str, Any] | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    day_filter = report_day or (output.stem[:10] if re.match(r"\d{4}-\d{2}-\d{2}", output.stem) else None)
    query = """SELECT t.task_id,t.created_at,t.sender_name,t.employee_name,
        t.original_filename,t.file_size,t.duration_seconds,t.accepted_number,t.status,r.score,r.qualified,
        r.summary,r.analysis,r.optimization,r.employee_speaker,r.speaker_confidence,r.dimension_scores,r.issues,
        r.needs_human_review,t.completed_at,o.call_summary,o.retention_result,o.uncertain_parts,o.full_dialogue,o.role_confidence,o.service_segments,o.customer_segments
        FROM audio_tasks t LEFT JOIN quality_results r ON r.task_id=t.task_id
        LEFT JOIN organized_calls o ON o.task_id=t.task_id
        ORDER BY t.rowid"""
    all_rows = connection.execute(query).fetchall()
    if day_filter is None:
        rows = all_rows
    else:
        rows = [
            row for row in all_rows
            if (extract_recording_date(str(row[4] or "")) or (str(row[1] or "")[:10])) == day_filter
        ]
    headers = [
        "序号",
        "日期",
        "员工姓名",
        "客户号码",
        "录音文件名",
        "通话时间",
        "问（25%）",
        "查（25%）",
        "比（25%）",
        "算（25%）",
        "问证据",
        "查证据",
        "比证据",
        "算证据",
        "是否有挽留动作",
        "挽留场景存在问题",
        "其他存在问题",
        "质检得分",
        "销降离原因",
        "挽留结果",
        "录音总结",
        "存在问题分析",
        "优化意见",
        "整理后的录音文本",
        "原始转录文本",
        "ASR不确定/需复核原因",
        "任务号",
        "质检人",
        "分公司",
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "降挽质检情况"
    sheet.append(headers)
    for index, raw in enumerate(rows, start=1):
        task_id, created_at, sender_name, employee_name, filename, file_size, duration_seconds, accepted_number, status, score, qualified, summary, analysis, optimization, employee_speaker, speaker_confidence, dimension_scores, issues, needs_review, completed_at, organized_summary, retention_result, uncertain_parts, full_dialogue, role_confidence, service_segments, customer_segments = raw
        recording_date = extract_recording_date(filename or "") or (created_at[:10] if created_at else "")
        rule_ids = _issue_rule_ids(issues)
        transcript_segments = load_segments(connection, task_id)
        transcript = join_transcript(transcript_segments)
        service_segment_list = _segments_from_json(service_segments, "客服")
        customer_segment_list = _segments_from_json(customer_segments, "客户")
        deepseek_actions = _load_deepseek_actions(connection, task_id)
        use_deepseek_actions = bool(deepseek_actions) and normalize_text(employee_speaker) == "DeepSeek语义审核"
        if use_deepseek_actions:
            action_evidence = {
                rule_id: filter_customer_perspective_evidence(
                    _deepseek_action_text(deepseek_actions or {}, rule_id),
                    customer_segment_list,
                )
                for rule_id in ("missing_ask", "missing_check", "missing_compare", "missing_calculate")
            }
            action_status = {
                rule_id: _weighted_action_status(rule_ids, rule_id)
                for rule_id in ("missing_ask", "missing_check", "missing_compare", "missing_calculate")
            }
        else:
            action_evidence = find_process_action_evidence(
                transcript_segments,
                None,
                transcript_segments,
                rules,
            )
            action_evidence = {
                rule_id: filter_customer_perspective_evidence(evidence, customer_segment_list)
                for rule_id, evidence in action_evidence.items()
            }
            action_status = {
                rule_id: _evidence_weighted_action_status(action_evidence[rule_id])
                for rule_id in ("missing_ask", "missing_check", "missing_compare", "missing_calculate")
            }
        service_reason = _infer_service_reason(customer_segment_list or transcript_segments, employee_speaker)
        review_reason = normalize_text(analysis) if needs_review else ""
        if needs_review and normalize_text(uncertain_parts):
            review_reason = (review_reason + "；" if review_reason else "") + f"DeepSeek不确定片段：{normalize_text(uncertain_parts)}"
        organized_text = full_dialogue or _segments_json_text(service_segments, "客服") or organized_summary or ""
        scene_rule_ids = {"missing_ask", "missing_check", "missing_calculate", "missing_compare", "missing_retention_action", "missing_retention_success"}
        other_problem = _compact_issues(issues, rules, scene_rule_ids)
        scene_problem = _retention_scene_problem(issues)
        sheet.append([
            index,
            recording_date,
            extract_employee_name(filename or "") or employee_name or sender_name or "",
            accepted_number or "",
            filename or "",
            completed_at or created_at or "",
            action_status["missing_ask"],
            action_status["missing_check"],
            action_status["missing_compare"],
            action_status["missing_calculate"],
            action_evidence["missing_ask"],
            action_evidence["missing_check"],
            action_evidence["missing_compare"],
            action_evidence["missing_calculate"],
            "否" if "missing_retention_action" in rule_ids else "是",
            scene_problem,
            other_problem,
            score if score is not None else "",
            service_reason,
            retention_result or "",
            organized_summary or summary or "",
            analysis or "",
            optimization or "",
            organized_text,
            transcript,
            review_reason,
            task_id,
            "",
            "",
        ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [8, 12, 14, 16, 34, 22, 14, 14, 14, 14, 42, 42, 42, 42, 18, 36, 36, 12, 20, 20, 48, 52, 52, 80, 90, 52, 24, 12, 12]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.stem + ".tmp" + output.suffix)
    workbook.save(temporary)
    try:
        temporary.replace(output)
        return output
    except PermissionError:
        fallback = output.with_name(f"{output.stem}-{datetime.now().strftime('%H%M%S')}{output.suffix}")
        temporary.replace(fallback)
        return fallback


def export_daily_report(connection: sqlite3.Connection, output_root: Path, when: datetime | None = None, rules: dict[str, Any] | None = None) -> Path:
    day = (when or datetime.now()).strftime("%Y-%m-%d")
    return export_report(connection, output_root / f"{day}_降挽质检情况.xlsx", day, rules)


def format_received_reply(task_id: str) -> str:
    return f"【听音质检】\n任务 {task_id} 已接收，正在处理。"


def format_completed_reply(task_id: str, needs_human_review: bool = False) -> str:
    return "听音检测完成，但需要人工复核。" if needs_human_review else "听音检测完成，已写入质检表。"
