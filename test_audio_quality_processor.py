import tempfile
import unittest
from pathlib import Path

from audio_quality.audio_quality_processor import (
    AudioQualityError, AudioSubmission, IncomingAudioMessage, TranscriptSegment,
    claim_next_task, deterministic_quality_check, evaluate_asr_quality, export_report,
    extract_accepted_number, extract_employee_name, extract_recording_date, format_completed_reply, init_db, load_json, mark_task_failed,
    filter_customer_perspective_evidence, parse_submission, parse_transcript, reserve_task, save_organized_call,
    save_deepseek_quality_assessment, save_quality_result, save_transcript, scan_rule_hits, find_process_action_evidence,
    run_daily_task_cleanup,
)
from audio_quality.batch_transcribe_faster_whisper import parse_filename_metadata
from audio_quality.deepseek_organizer import OrganizedCall
from audio_quality.deepseek_quality import assessment_to_quality_result, reconcile_retention_semantics, reinforce_ask_evidence, reinforce_location_lookup_evidence, reinforce_observed_record_evidence, reinforce_verification_lookup_evidence, validate_assessment


RULES = {
    "trigger": "#听音检测",
    "trigger_aliases": ["#听音检测", "#听音质检"],
    "order_id_pattern": r"[A-Za-z0-9_-]{4,64}",
    "supported_audio_suffixes": [".mp3", ".wav"],
    "max_file_size_bytes": 1000,
    "qualified_score": 80,
    "speaker_review_threshold": 0.75,
    "employee_speaker_keywords": ["中国移动", "您好"],
    "dimensions": {
        "basic_courtesy": {"weight": 50, "base_score": 100, "rule_ids": ["P01"]},
        "respectful_expression": {"weight": 50, "base_score": 100, "rule_ids": ["P01"]},
    },
    "phrase_rules": [{"id": "P01", "pattern": "我也没有办法", "description": "消极表达", "deduction": 10}],
}


def make_message(**overrides):
    values = {
        "message_id": "m1", "group_id": "g1", "sender_user_id": "u1", "sender_name": "提交人",
        "text": "@Bot #听音检测\n员工：张三\n工单号：A1234\n受理号码：15978157631",
        "attachment_id": "f1", "filename": "call.mp3", "file_size": 500,
        "duration_seconds": 30, "attachment_hash": "hash1",
    }
    values.update(overrides)
    return IncomingAudioMessage(**values)


class AudioQualityProcessorTests(unittest.TestCase):
    def test_deepseek_assessment_requires_verbatim_transcript_evidence(self):
        transcript = "[0.0-3.0] CH0: 您为什么想取消宽带，是不是费用高？"
        assessment = validate_assessment({
            "actions": {
                "missing_ask": {"present": True, "evidence": ["您为什么想取消宽带", "模型编造的证据"]},
            },
            "summary": "测试",
        }, transcript)
        self.assertTrue(assessment.actions["missing_ask"]["present"])
        self.assertEqual(["您为什么想取消宽带"], assessment.actions["missing_ask"]["evidence"])
        self.assertFalse(assessment.actions["missing_check"]["present"])

    def test_deepseek_question_guard_keeps_complaint_reason_question(self):
        segments = [
            TranscriptSegment(2.3, 4.5, "CH0", "我是移动的工作人员。"),
            TranscriptSegment(4.5, 10.0, "CH0", "就是你投诉了一些，就是说问题是什么原因呢？"),
        ]
        assessment = validate_assessment({"actions": {}, "summary": ""}, "\n".join(item.text for item in segments))
        repaired = reinforce_ask_evidence(assessment, segments)
        self.assertTrue(repaired.actions["missing_ask"]["present"])
        self.assertIn("问题是什么原因", repaired.actions["missing_ask"]["evidence"][0])

    def test_deepseek_question_guard_accepts_work_order_confirmation(self):
        segments = [
            TranscriptSegment(
                0.6,
                12.5,
                "CH0",
                "喂你好，我这里是移动后台处理投诉专员，我这里看到你有一个工单是投诉宽带问题的是吗？",
            )
        ]
        assessment = validate_assessment({"actions": {}, "summary": ""}, segments[0].text)
        repaired = reinforce_ask_evidence(assessment, segments)
        repaired = reinforce_observed_record_evidence(repaired, segments)
        self.assertTrue(repaired.actions["missing_ask"]["present"])
        self.assertTrue(repaired.actions["missing_check"]["present"])
        self.assertIn("投诉宽带问题的是吗", repaired.actions["missing_ask"]["evidence"][0])
        self.assertIn("看到你有一个工单", repaired.actions["missing_check"]["evidence"][0])

    def test_deepseek_question_guard_does_not_treat_later_customer_question_as_staff_evidence(self):
        segments = [
            TranscriptSegment(2.3, 4.5, "CH0", "我是移动的工作人员。"),
            TranscriptSegment(4.5, 10.0, "CH0", "你投诉的宽带问题是什么原因呢？"),
            TranscriptSegment(20.0, 25.0, "CH0", "那我去哪个营业厅？"),
        ]
        assessment = validate_assessment({"actions": {}, "summary": ""}, "\n".join(item.text for item in segments))
        repaired = reinforce_ask_evidence(assessment, segments)
        self.assertEqual(["你投诉的宽带问题是什么原因呢?"], repaired.actions["missing_ask"]["evidence"])

    def test_all_relevant_action_evidence_is_preserved(self):
        quotes = [f"证据{i}" for i in range(1, 5)]
        transcript = "；".join(quotes)
        assessment = validate_assessment({
            "actions": {
                "missing_ask": {"present": True, "evidence": quotes},
                "missing_check": {"present": True, "evidence": quotes},
                "missing_compare": {"present": True, "evidence": quotes},
            },
            "summary": "",
        }, transcript)
        self.assertEqual(quotes, assessment.actions["missing_ask"]["evidence"])
        self.assertEqual(quotes, assessment.actions["missing_check"]["evidence"])
        self.assertEqual(quotes, assessment.actions["missing_compare"]["evidence"])

    def test_address_to_office_lookup_is_check_evidence(self):
        segments = [
            TranscriptSegment(0, 3, "CH0", "您这边的小区和地址是哪里呢？"),
            TranscriptSegment(4, 6, "CH0", "我住在春风社区。"),
            TranscriptSegment(7, 10, "CH0", "可以去中心营业厅，该厅有权限办理。"),
            TranscriptSegment(11, 13, "CH0", "那我去哪个营业厅？"),
        ]
        assessment = validate_assessment({"actions": {}, "summary": ""}, "\n".join(item.text for item in segments))
        repaired = reinforce_location_lookup_evidence(assessment, segments)
        self.assertTrue(repaired.actions["missing_check"]["present"])
        self.assertEqual(["可以去中心营业厅,该厅有权限办理。"], repaired.actions["missing_check"]["evidence"])

    def test_system_observation_is_check_evidence(self):
        segments = [TranscriptSegment(0, 8, "CH0", "我从系统看到您当前有一个宽带投诉工单。")]
        assessment = validate_assessment({"actions": {}, "summary": ""}, segments[0].text)
        repaired = reinforce_observed_record_evidence(assessment, segments)
        self.assertTrue(repaired.actions["missing_check"]["present"])
        self.assertEqual(["我从系统看到您当前有一个宽带投诉工单。"], repaired.actions["missing_check"]["evidence"])

    def test_verification_code_lookup_reinforces_check_and_calculation(self):
        segments = [
            TranscriptSegment(252.3, 255.1, "CH0", "您方便帮我看个短信验证码吗？"),
            TranscriptSegment(255.9, 260.8, "CH0", "稍等一下，我找到了，就是那增值业务费28块钱。"),
            TranscriptSegment(261.0, 266.0, "CH0", "有的时候您是在中国移动APP看的是吗？"),
            TranscriptSegment(267.0, 275.0, "CH0", "他这里有一个移动高清基础包，还有一个10元生活权益包。"),
            TranscriptSegment(276.0, 280.0, "CH0", "那个10元生活权益包是可以去领话费券的。"),
        ]
        transcript = "\n".join(item.text for item in segments)
        assessment = validate_assessment({"actions": {}, "summary": ""}, transcript)
        repaired = reinforce_verification_lookup_evidence(assessment, segments)
        self.assertTrue(repaired.actions["missing_check"]["present"])
        self.assertTrue(repaired.actions["missing_calculate"]["present"])
        self.assertIn("短信验证码", " ".join(repaired.actions["missing_check"]["evidence"]))
        self.assertIn("增值业务费28块钱", " ".join(repaired.actions["missing_calculate"]["evidence"]))
        self.assertIn("话费券", " ".join(repaired.actions["missing_calculate"]["evidence"]))

    def test_discount_offer_links_compare_calculate_and_retention(self):
        segments = [
            TranscriptSegment(0, 4, "CH0", "原来是每月28块钱。"),
            TranscriptSegment(5, 10, "CH0", "我们这边帮您申请优惠一年,费用都给您免掉,您看可以吗？"),
        ]
        assessment = validate_assessment({"actions": {}, "summary": ""}, "\n".join(item.text for item in segments))
        repaired = reconcile_retention_semantics(assessment, segments)
        self.assertTrue(repaired.actions["missing_compare"]["present"])
        self.assertTrue(repaired.actions["missing_calculate"]["present"])
        self.assertTrue(repaired.actions["missing_retention_action"]["present"])

    def test_deepseek_retention_guard_accepts_discount_and_zero_yuan_offer(self):
        segments = [
            TranscriptSegment(0, 5, "CH0", "那我这边帮你申请做个优惠，你还是在88套餐里算了。"),
            TranscriptSegment(6, 11, "CH0", "我可以给你这个宽带改0元的，不用钱的，可以继续使用。"),
        ]
        assessment = validate_assessment({"actions": {}, "summary": ""}, "\n".join(item.text for item in segments))
        repaired = reconcile_retention_semantics(assessment, segments)
        self.assertTrue(repaired.actions["missing_retention_action"]["present"])
        self.assertIn("申请做个优惠", " ".join(repaired.actions["missing_retention_action"]["evidence"]))

    def test_positive_acceptance_reason_cannot_remain_failed_retention(self):
        segments = [TranscriptSegment(0, 8, "CH0", "我这边帮你申请做个优惠,嗯好的那就帮我申请吧。")]
        assessment = validate_assessment({
            "actions": {
                "missing_retention_success": {
                    "present": False,
                    "evidence": [],
                    "reason": "客户未逐字说同意。",
                },
            },
            "summary": "",
        }, segments[0].text)
        repaired = reconcile_retention_semantics(assessment, segments)
        self.assertTrue(repaired.actions["missing_retention_success"]["present"])

    def test_rejection_after_offer_keeps_retention_unsuccessful(self):
        segments = [
            TranscriptSegment(0, 5, "CH0", "我们帮您申请优惠一年,可以吗？"),
            TranscriptSegment(6, 9, "CH0", "那我不用了,还是要退掉。"),
        ]
        assessment = validate_assessment({"actions": {}, "summary": ""}, "\n".join(item.text for item in segments))
        repaired = reconcile_retention_semantics(assessment, segments)
        self.assertTrue(repaired.actions["missing_retention_action"]["present"])
        self.assertFalse(repaired.actions["missing_retention_success"]["present"])

    def test_parse_audio_filename_metadata(self):
        self.assertEqual(("2026-07-23", "15978157631"), parse_filename_metadata("2026-7-23 雷建宏 15978157631号码.m4a"))

    def test_parse_submission_and_attachment(self):
        self.assertEqual(AudioSubmission("张三", "A1234", "15978157631"), parse_submission(make_message().text, RULES))
        with self.assertRaisesRegex(AudioQualityError, "只能执行一个功能"):
            parse_submission("#听音检测 #沟通记录\n员工：张三", RULES)

    def test_extract_accepted_number(self):
        self.assertEqual("15978157631", extract_accepted_number("2026-7-23 15978157631号码.m4a"))

    def test_extract_employee_name_from_grid_filename(self):
        self.assertEqual("黄荣宇", extract_employee_name("2026年7月10号南阳网格黄荣宇13978877634+移动高清.mp3"))
        self.assertEqual("韦振珮", extract_employee_name("2026年7月10号南阳网格韦振珮+13481118432+移动高清(1).m4a"))
        self.assertEqual("张三", extract_employee_name("@Bot #听音检测\n员工：张三"))

    def test_action_evidence_accepts_strong_content_from_uncertain_speaker(self):
        rules = load_json(Path("audio_quality/audio_quality_rules.json"))
        raw_segments = [
            TranscriptSegment(70.8, 90.4, "CH0", "我改套餐是比如说你们现在这个套餐是你们这个宽带，宽带这个套餐就呢卡在这的话我就改不了，八号的8块钱的八号套餐你明白吗？如果是续约的话可以改成1元套餐，就是1元的宽带。"),
            TranscriptSegment(93.2, 95.0, "CH0", "问题是我不需要宽带了呀。"),
        ]
        evidence = find_process_action_evidence(
            [TranscriptSegment(93.2, 95.0, "客服", "问题是我不需要宽带了呀。")],
            "客服",
            raw_segments,
            rules,
        )
        self.assertNotEqual("未找到明确证据", evidence["missing_compare"])
        self.assertNotEqual("未找到明确证据", evidence["missing_calculate"])

    def test_action_evidence_uses_content_when_speaker_labels_are_wrong(self):
        rules = load_json(Path("audio_quality/audio_quality_rules.json"))
        segments = [
            TranscriptSegment(0, 6, "客户", "如果是这样的话，我继续给你做优惠两年可以吗？因为我们要帮你申请免违约金，所以必须是那几个营业厅才能帮你申请得了违约金。"),
            TranscriptSegment(7, 9, "客服", "那我去哪个营业厅？"),
            TranscriptSegment(10, 16, "客户", "现在只有两个方案，一个是你去回收设备后我再帮你补退费用，另一个是我这边直接上单帮你补退费用。"),
        ]
        evidence = find_process_action_evidence([segments[1]], "客服", segments, rules)
        self.assertIn("优惠两年", evidence["missing_compare"])
        self.assertIn("两个方案", evidence["missing_compare"])
        self.assertIn("违约金", evidence["missing_calculate"])
        self.assertNotIn("那我去哪个营业厅", evidence["missing_ask"])

    def test_action_evidence_trims_long_mixed_customer_reply(self):
        rules = load_json(Path("audio_quality/audio_quality_rules.json"))
        long_text = (
            "因为我不需要这个机顶盒了，你们为什么还收费，我肯定是要投诉的。"
            "这个问题我已经说了很多遍，我就是客户自己这边不想用了，也不想再跑来跑去。"
            "你们前面问我什么时候方便，我也只是回答我没时间去营业厅。"
            "如果是这样的话，我继续给你做优惠两年可以吗？因为这个我们要帮你申请免违约金，所以你必须去指定营业厅。"
            "我出时间去退这个设备我可以了，我认了，但是你们还要给我补退。"
            "现在只有两个方案，一个回收以后补退，一个我这边直接上单补退。"
        )
        evidence = find_process_action_evidence([], "客服", [TranscriptSegment(0, 40, "CH0", long_text)], rules)
        self.assertIn("优惠两年", evidence["missing_compare"])
        self.assertLess(len(evidence["missing_compare"]), len(long_text))
        self.assertNotIn("我肯定是要投诉", evidence["missing_compare"])

    def test_discount_and_monthly_add_on_are_business_evidence(self):
        rules = load_json(Path("audio_quality/audio_quality_rules.json"))
        segments = [
            TranscriptSegment(46.1, 63.9, "CH0", "我们这边可以把这个资费优惠给你，优惠继续使用一年，就是给你优惠15块。"),
            TranscriptSegment(47.7, 68.7, "CH0", "这个可以每个月加1块钱，你就可以继续使用那个宽带，然后也可以继续改套餐。"),
            TranscriptSegment(67.3, 78.2, "CH0", "我就帮你改成0元的宽带也可以，就是不用钱的，0元宽带一直都可以免费使用。"),
        ]
        evidence = find_process_action_evidence([], "客服", segments, rules)
        self.assertIn("优惠15块", evidence["missing_compare"])
        self.assertIn("每个月加1块钱", evidence["missing_calculate"])
        result = deterministic_quality_check(
            segments,
            rules,
            employee_speaker_override="CH0",
            speaker_confidence_override=0.9,
            asr_quality_segments=segments,
        )
        self.assertNotIn("missing_retention_action", {issue.rule_id for issue in result.issues})

    def test_plan_fee_breakdown_is_calculation(self):
        rules = load_json(Path("audio_quality/audio_quality_rules.json"))
        segments = [TranscriptSegment(
            255.9,
            295.2,
            "CH0",
            "我找到了，就是增值业务费28块。他有时候是这个价钱，也不一定是机顶盒收费。他这里有一个移动高清基础包，还有一个10元生活权益包，10元生活权益包可以领话费券。",
        )]
        evidence = find_process_action_evidence([], "客服", segments, rules)
        self.assertIn("增值业务费28块", evidence["missing_calculate"])
        self.assertIn("话费券", evidence["missing_calculate"])

    def test_plan_price_change_and_two_options_are_comparison_evidence(self):
        rules = load_json(Path("audio_quality/audio_quality_rules.json"))
        segments = [
            TranscriptSegment(0, 10, "CH0", "我88元套餐以前一直用，正常可以改成78元套餐，88元套餐包含机顶盒。"),
            TranscriptSegment(10, 20, "CH0", "有两个处理方案，一个去营业厅补退费用，或者我上单帮你补退费用。"),
        ]
        evidence = find_process_action_evidence([], "客服", segments, rules)
        self.assertIn("改成78元套餐", evidence["missing_compare"])
        self.assertIn("两个处理方案", evidence["missing_compare"])
        self.assertIn("改成78元套餐", evidence["missing_calculate"])

    def test_spoken_discount_with_repeated_amount_is_comparison(self):
        rules = load_json(Path("audio_quality/audio_quality_rules.json"))
        segments = [TranscriptSegment(43.5, 49.3, "CH0", "优惠一年，就是那个28、20多块钱的，帮您申请优惠。")]
        evidence = find_process_action_evidence([], "客服", segments, rules)
        self.assertIn("20多块钱", evidence["missing_compare"])

    def test_customer_acceptance_overrides_pending_retention_result(self):
        rules = load_json(Path("audio_quality/audio_quality_rules.json"))
        segments = [TranscriptSegment(
            83.6,
            114.8,
            "客服",
            "客服说可以优惠一年，客户说那我先试一下吧。",
        )]
        result = deterministic_quality_check(
            segments,
            rules,
            employee_speaker_override="客服",
            speaker_confidence_override=0.9,
            asr_quality_segments=segments,
            retention_result="待跟进",
        )
        self.assertNotIn("missing_retention_success", {issue.rule_id for issue in result.issues})

    def test_check_evidence_can_span_adjacent_asr_segments(self):
        rules = load_json(Path("audio_quality/audio_quality_rules.json"))
        segments = [
            TranscriptSegment(243.9, 245.8, "CH0", "我帮你查一下。"),
            TranscriptSegment(252.3, 255.1, "CH0", "您方便看个短信验证码吗？"),
            TranscriptSegment(255.9, 260.8, "CH0", "稍等，我找到了，就是增值业务费28块钱。"),
        ]
        evidence = find_process_action_evidence([], "客服", segments, rules)
        self.assertIn("验证码", evidence["missing_check"])
        self.assertIn("增值业务费28块钱", evidence["missing_check"])

    def test_complaint_work_order_confirmation_counts_as_ask_and_check(self):
        rules = load_json(Path("audio_quality/audio_quality_rules.json"))
        segments = [
            TranscriptSegment(
                0.6,
                12.5,
                "CH0",
                "喂你好，我这里是移动后台处理投诉专员，我这里看到你有一个工单是投诉宽带问题的是吗？",
            ),
            TranscriptSegment(
                15.0,
                22.4,
                "CH0",
                "投诉宽带问题啊，对你之前打过10086投诉宽带关于他宽带的问题吗？",
            ),
        ]
        evidence = find_process_action_evidence([], "客服", segments, rules)
        self.assertIn("投诉宽带问题的是吗", evidence["missing_ask"])
        self.assertIn("看到你有一个工单", evidence["missing_check"])

    def test_empty_semantic_match_falls_back_to_configured_rules(self):
        rules = load_json(Path("audio_quality/audio_quality_rules.json"))
        segments = [
            TranscriptSegment(125.4, 134.4, "CH0", "您可以去附近的桂春厅或者是云圆湖厅那里去取消哦！"),
            TranscriptSegment(470.6, 485.3, "CH0", "需要看你们那一边的营业厅有多少套餐给您，有时候他那边的权限不一定有。"),
        ]
        evidence = find_process_action_evidence([], "客服", segments, rules)
        self.assertIn("桂春厅", evidence["missing_check"])
        self.assertIn("权限不一定有", evidence["missing_check"])

    def test_output_filter_removes_evidence_overlapping_organized_customer_segments(self):
        evidence = (
            "[1.0-2.0] CH0：我们人不在南宁，你们有什么办法可以取消？\n"
            "[3.0-4.0] CH0：您可以去附近营业厅那里取消哦。"
        )
        customer_segments = [
            TranscriptSegment(1.0, 2.0, "客户", "我们人不在南宁，你们有什么办法可以取消？"),
        ]
        filtered = filter_customer_perspective_evidence(evidence, customer_segments)
        self.assertNotIn("我们人不在南宁", filtered)
        self.assertIn("附近营业厅", filtered)

    def test_output_filter_does_not_trust_speaker_label_without_text_overlap(self):
        evidence = "[3.0-4.0] 客户：您可以去附近营业厅那里取消哦。"
        customer_segments = [
            TranscriptSegment(1.0, 2.0, "客户", "我们人不在南宁，你们有什么办法可以取消？"),
        ]
        filtered = filter_customer_perspective_evidence(evidence, customer_segments)
        self.assertIn("附近营业厅", filtered)

    def test_retention_action_accepts_apply_discount_and_zero_yuan_offer(self):
        rules = load_json(Path("audio_quality/audio_quality_rules.json"))
        segments = [
            TranscriptSegment(
                0,
                5,
                "客服",
                "那我这边帮你申请做个优惠，你还是在88套餐里算了，我先帮你申请上去做优惠。",
            ),
            TranscriptSegment(
                6,
                11,
                "客服",
                "那我可以给你这个宽带改0元的，不用钱的，可以继续使用，可以吗？",
            ),
        ]
        result = deterministic_quality_check(
            segments,
            rules,
            employee_speaker_override="客服",
            speaker_confidence_override=0.9,
            asr_quality_segments=segments,
        )
        self.assertNotIn("missing_retention_action", {issue.rule_id for issue in result.issues})

    def test_semantic_action_definitions_require_meaningful_evidence(self):
        rules = load_json(Path("audio_quality/audio_quality_rules.json"))
        segments = [
            TranscriptSegment(0, 2, "CH0", "您为什么想把套餐改掉，是觉得费用高还是用量不够？"),
            TranscriptSegment(3, 5, "CH0", "我帮您查到当前是88元套餐，近三个月流量只用了2G。"),
            TranscriptSegment(6, 8, "CH0", "我建议您改成58元套餐，比现在每月省30元，还有更多流量。"),
        ]
        evidence = find_process_action_evidence([], "客服", segments, rules)
        self.assertIn("为什么想", evidence["missing_ask"])
        self.assertIn("查到当前是88元套餐", evidence["missing_check"])
        self.assertIn("改成58元套餐", evidence["missing_compare"])
        self.assertIn("每月省30元", evidence["missing_calculate"])

    def test_extract_recording_date_from_filename(self):
        self.assertEqual("2026-07-10", extract_recording_date("2026年7月10号南阳网格黄荣宇13978877634+移动高清.mp3"))
        self.assertEqual("2026-07-23", extract_recording_date("2026-7-23 雷建宏 15978157631号码.m4a"))

    def test_transcript_validation_and_rule_scan(self):
        segments = parse_transcript([
            {"start_seconds": 0, "end_seconds": 2, "speaker": "EMP", "text": "中国移动您好"},
            {"start_seconds": 3, "end_seconds": 5, "speaker": "EMP", "text": "这个我也没有办法"},
        ])
        self.assertEqual(2, len(segments))
        self.assertEqual(1, len(scan_rule_hits(segments, "EMP", RULES)))
        result = deterministic_quality_check(segments, RULES)
        self.assertEqual(90, result.score)
        self.assertEqual(1, len(result.issues))

    def test_asr_quality_gate_requires_review(self):
        rules = dict(RULES)
        rules["asr_quality"] = {"enabled": True, "min_total_characters": 20, "min_segment_count": 2, "key_business_terms": ["套餐"]}
        segments = [TranscriptSegment(0, 2, "EMP", "中国移动您好")]
        self.assertTrue(evaluate_asr_quality(segments, rules))
        result = deterministic_quality_check(segments, rules)
        self.assertTrue(result.needs_human_review)
        self.assertFalse(result.qualified)

    def test_asr_quality_does_not_require_two_speakers_by_default(self):
        rules = dict(RULES)
        rules["asr_quality"] = {
            "enabled": True,
            "min_total_characters": 4,
            "min_segment_count": 1,
            "min_distinct_speakers": 2,
            "key_business_terms": ["套餐"],
        }
        segments = [TranscriptSegment(0, 2, "CH0", "客户咨询套餐费用，希望办理优惠套餐")]
        self.assertEqual([], evaluate_asr_quality(segments, rules))

    def test_asr_quality_can_require_speaker_separation_when_configured(self):
        rules = dict(RULES)
        rules["asr_quality"] = {
            "enabled": True,
            "require_speaker_separation": True,
            "min_total_characters": 4,
            "min_segment_count": 1,
            "min_distinct_speakers": 2,
            "key_business_terms": ["套餐"],
        }
        segments = [TranscriptSegment(0, 2, "CH0", "客户咨询套餐费用，希望办理优惠套餐")]
        self.assertEqual(["ASR可用说话人标签少于2个"], evaluate_asr_quality(segments, rules))

    def test_process_action_evidence_is_observable_without_deductions(self):
        rules = dict(RULES)
        rules["required_action_scoring"] = {"evidence_only": True}
        rules["score_policy"] = {"mode": "flat_deduction", "base_score": 100}
        rules["required_actions"] = [
            {"id": action, "patterns": ["never-match"], "deduction": 25}
            for action in ("missing_ask", "missing_check", "missing_compare", "missing_calculate")
        ]
        segments = [TranscriptSegment(
            0, 15, "客服",
            "您这次想办理什么业务，是什么原因呢？我帮您查一下当前套餐和流量使用情况。"
            "原套餐和现在套餐相比，18元套餐每月能省10元。",
        )]
        evidence = find_process_action_evidence(segments, "客服")
        self.assertTrue(all(evidence.values()))
        result = deterministic_quality_check(segments, rules, employee_speaker_override="客服", speaker_confidence_override=0.9)
        self.assertEqual(100, result.score)
        self.assertEqual([], result.issues)

    def test_phrase_rule_exclusions_allow_business_condition_explanations(self):
        rules = {
            "speaker_review_threshold": 0.5,
            "employee_speaker_keywords": ["申请", "违约金"],
            "score_policy": {"mode": "flat_deduction", "base_score": 100},
            "phrase_rules": [
                {
                    "id": "commanding_tone",
                    "pattern": "你必须|你赶紧|你快点|别再|你只能",
                    "exclude_patterns": [
                        "必须.{0,20}(才能|才可以|才可|方可).{0,20}(申请|办理|取消|处理|免|减免|退费|补退|违约金)"
                    ],
                    "regex": True,
                    "description": "疑似命令式语气",
                    "deduction": 12,
                }
            ],
        }
        segments = [
            TranscriptSegment(0, 5, "客服", "因为我们要帮你申请免违约金，所以你必须要是那几个营业厅才能帮你申请得了违约金。")
        ]
        result = deterministic_quality_check(segments, rules, employee_speaker_override="客服", speaker_confidence_override=0.9)
        self.assertEqual([], result.issues)
        self.assertEqual(100, result.score)

    def test_analysis_uses_chinese_rule_labels(self):
        rules = dict(RULES)
        rules["score_policy"] = {"mode": "flat_deduction", "base_score": 100}
        rules["required_action_scoring"] = {"evidence_only": False}
        rules["action_evidence_rules"] = {"missing_check": {"patterns": ["不会命中"]}}
        rules["required_actions"] = [
            {"id": "missing_check", "label": "查", "description": "未找到查套餐/查用量/查限制的明确证据", "deduction": 5}
        ]
        result = deterministic_quality_check(
            [TranscriptSegment(0, 2, "客服", "中国移动您好。")],
            rules,
            employee_speaker_override="客服",
            speaker_confidence_override=0.9,
        )
        self.assertIn("查：未找到查套餐/查用量/查限制的明确证据", result.analysis)
        self.assertNotIn("missing_check", result.analysis)

    def test_database_single_sheet_report(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            connection = init_db(Path(directory) / "state.sqlite3")
            task_id, created = reserve_task(connection, make_message(), parse_submission(make_message().text, RULES))
            self.assertTrue(created)
            self.assertEqual(task_id, claim_next_task(connection, "worker-1"))
            mark_task_failed(connection, task_id, "asr", "temporary")
            row = connection.execute("SELECT status,error_stage FROM audio_tasks WHERE task_id=?", (task_id,)).fetchone()
            self.assertEqual(("failed", "asr"), row)
            connection.execute("UPDATE audio_tasks SET status='received',error_stage=NULL,error_message=NULL WHERE task_id=?", (task_id,))
            connection.commit()
            self.assertEqual(task_id, claim_next_task(connection, "worker-2"))
            segments = [TranscriptSegment(0, 2, "客服", "中国移动您好，我帮您查询套餐。")]
            save_transcript(connection, task_id, segments)
            customer_segments = [TranscriptSegment(3, 5, "客户", "套餐太贵了，我想改8元保号。")]
            organized = OrganizedCall(segments, customer_segments, "客服：中国移动您好，我帮您查询套餐。\n客户：套餐太贵了，我想改8元保号。", "客户咨询套餐", "引导", "", 0.9)
            save_organized_call(connection, task_id, organized)
            result = deterministic_quality_check(segments, RULES, employee_speaker_override="客服", speaker_confidence_override=0.9, asr_quality_segments=segments)
            save_quality_result(connection, task_id, result)
            output = export_report(connection, Path(directory) / "report.xlsx")
            connection.close()
            sheet = load_workbook(output).active
            headers = [cell.value for cell in sheet[1]]
            self.assertIn("录音文件名", headers)
            self.assertIn("问证据", headers)
            self.assertIn("查证据", headers)
            self.assertIn("比证据", headers)
            self.assertIn("算证据", headers)
            self.assertIn("挽留场景存在问题", headers)
            self.assertIn("整理后的录音文本", headers)
            self.assertIn("原始转录文本", headers)
            self.assertIn("任务号", headers)
            self.assertNotIn("原始转录文本文件", headers)
            transcript = sheet.cell(row=2, column=headers.index("原始转录文本") + 1).value
            self.assertIn("客服", transcript)
            service_reason = sheet.cell(row=2, column=headers.index("销降离原因") + 1).value
            self.assertIn("偏高", service_reason)

    def test_report_issue_text_uses_chinese_labels(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            connection = init_db(Path(directory) / "state.sqlite3")
            task_id, created = reserve_task(connection, make_message(), parse_submission(make_message().text, RULES))
            self.assertTrue(created)
            segments = [TranscriptSegment(0, 2, "客服", "中国移动您好。")]
            save_transcript(connection, task_id, segments)
            rules = dict(RULES)
            rules["score_policy"] = {"mode": "flat_deduction", "base_score": 100}
            rules["required_action_scoring"] = {"evidence_only": False}
            rules["action_evidence_rules"] = {"missing_check": {"patterns": ["不会命中"]}}
            rules["required_actions"] = [
                {"id": "missing_check", "label": "查", "description": "未找到查套餐/查用量/查限制的明确证据", "deduction": 5}
            ]
            result = deterministic_quality_check(
                segments,
                rules,
                employee_speaker_override="客服",
                speaker_confidence_override=0.9,
            )
            save_quality_result(connection, task_id, result)
            output = export_report(connection, Path(directory) / "report.xlsx", rules=rules)
            connection.close()
            sheet = load_workbook(output).active
            headers = [cell.value for cell in sheet[1]]
            scene_problem = sheet.cell(row=2, column=headers.index("挽留场景存在问题") + 1).value
            other_problem = sheet.cell(row=2, column=headers.index("其他存在问题") + 1).value
            self.assertIn("未找到查套餐/查用量/查限制的明确证据", scene_problem)
            self.assertNotIn("missing_check", scene_problem)
            self.assertFalse(other_problem)

    def test_report_uses_deepseek_actions_for_primary_semantic_result(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            connection = init_db(Path(directory) / "state.sqlite3")
            task_id, created = reserve_task(connection, make_message(), parse_submission(make_message().text, RULES))
            self.assertTrue(created)
            segments = [
                TranscriptSegment(0, 3, "CH0", "您为什么想取消宽带，是不是费用高？"),
                TranscriptSegment(4, 8, "CH0", "那我不用了，还是要退掉。"),
            ]
            save_transcript(connection, task_id, segments)
            assessment = validate_assessment({
                "actions": {
                    "missing_ask": {
                        "present": True,
                        "evidence": ["您为什么想取消宽带"],
                        "reason": "客服询问取消原因。",
                    },
                    "missing_check": {
                        "present": False,
                        "evidence": [],
                        "reason": "客服未查询并告知客户实际情况。",
                    },
                },
                "summary": "客户要求取消宽带",
            }, "\n".join(item.text for item in segments))
            result = assessment_to_quality_result(assessment, segments, RULES)
            save_deepseek_quality_assessment(connection, task_id, result, assessment)
            connection.commit()
            save_quality_result(connection, task_id, result)
            output = export_report(connection, Path(directory) / "report.xlsx", rules=RULES)
            connection.close()
            sheet = load_workbook(output).active
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual("25%（已体现）", sheet.cell(row=2, column=headers.index("问（25%）") + 1).value)
            self.assertEqual("0%（未体现）", sheet.cell(row=2, column=headers.index("查（25%）") + 1).value)
            self.assertIn("已体现：您为什么想取消宽带", sheet.cell(row=2, column=headers.index("问证据") + 1).value)
            self.assertIn("未体现：客服未查询", sheet.cell(row=2, column=headers.index("查证据") + 1).value)

    def test_report_uses_recording_date_from_filename(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            connection = init_db(Path(directory) / "state.sqlite3")
            message = make_message(filename="2026年7月10号南阳网格黄荣宇13978877634+移动高清.mp3")
            task_id, created = reserve_task(connection, message, AudioSubmission("黄荣宇", None, "13978877634"))
            self.assertTrue(created)
            save_transcript(connection, task_id, [TranscriptSegment(0, 2, "客服", "中国移动您好。")])
            save_quality_result(connection, task_id, deterministic_quality_check(
                [TranscriptSegment(0, 2, "客服", "中国移动您好。")],
                RULES,
                employee_speaker_override="客服",
                speaker_confidence_override=0.9,
            ))
            output = export_report(connection, Path(directory) / "2026-07-10_降挽质检情况.xlsx", "2026-07-10")
            empty_output = export_report(connection, Path(directory) / "2026-07-29_降挽质检情况.xlsx", "2026-07-29")
            connection.close()
            sheet = load_workbook(output).active
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual("2026-07-10", sheet.cell(row=2, column=headers.index("日期") + 1).value)
            self.assertEqual(1, load_workbook(empty_output).active.max_row)

    def test_claim_next_task_handles_multiple_queue_items_once_each(self):
        with tempfile.TemporaryDirectory() as directory:
            connection = init_db(Path(directory) / "state.sqlite3")
            first, _ = reserve_task(connection, make_message(message_id="m-1", attachment_id="a-1", attachment_hash="h-1"), parse_submission(make_message().text, RULES))
            second, _ = reserve_task(connection, make_message(message_id="m-2", attachment_id="a-2", attachment_hash="h-2"), parse_submission(make_message().text, RULES))
            self.assertEqual(first, claim_next_task(connection, "worker-1"))
            self.assertEqual(second, claim_next_task(connection, "worker-1"))
            self.assertIsNone(claim_next_task(connection, "worker-1"))
            rows = connection.execute("SELECT task_id,status FROM audio_tasks ORDER BY rowid").fetchall()
            connection.close()
            self.assertEqual([(first, "processing"), (second, "processing")], rows)

    def test_daily_cleanup_preserves_pending_tasks(self):
        with tempfile.TemporaryDirectory() as directory:
            connection = init_db(Path(directory) / "state.sqlite3")
            task_id, created = reserve_task(connection, make_message(), parse_submission(make_message().text, RULES))
            self.assertTrue(created)
            self.assertTrue(run_daily_task_cleanup(connection, {"enabled": True}, today="2026-07-28"))
            self.assertEqual(1, connection.execute("SELECT count(*) FROM audio_tasks").fetchone()[0])
            reserve_task(connection, make_message(message_id="m-2", attachment_id="a-2", attachment_hash="h-2"), parse_submission(make_message().text, RULES))
            self.assertFalse(run_daily_task_cleanup(connection, {"enabled": True}, today="2026-07-28"))
            self.assertEqual(2, connection.execute("SELECT count(*) FROM audio_tasks").fetchone()[0])
            self.assertTrue(run_daily_task_cleanup(connection, {"enabled": True}, today="2026-07-29"))
            self.assertEqual(2, connection.execute("SELECT count(*) FROM audio_tasks").fetchone()[0])
            connection.close()

    def test_daily_cleanup_clear_all_requires_explicit_opt_in(self):
        with tempfile.TemporaryDirectory() as directory:
            connection = init_db(Path(directory) / "state.sqlite3")
            _, created = reserve_task(connection, make_message(), parse_submission(make_message().text, RULES))
            self.assertTrue(created)
            self.assertTrue(run_daily_task_cleanup(connection, {"enabled": True, "clear_all": True}, today="2026-07-28"))
            self.assertEqual(0, connection.execute("SELECT count(*) FROM audio_tasks").fetchone()[0])
            connection.close()

    def test_completed_reply(self):
        self.assertEqual("听音检测完成，已写入质检表。", format_completed_reply("QA-1"))
        self.assertEqual("听音检测完成，但需要人工复核。", format_completed_reply("QA-1", True))


if __name__ == "__main__":
    unittest.main()
