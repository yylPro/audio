from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable


SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".csv"}
SCIENTIFIC_NOTATION_RE = re.compile(r"^[+-]?\d+(?:\.\d+)?[eE][+-]?\d+$")
MENTION_RE = re.compile(r"(?:^|\s)@(\S+?)(?=\s|$|[，。,.：:])")


@dataclass(frozen=True)
class ImportResult:
    source: Path
    digest: str
    total_rows: int
    pending_rows: int
    counts: Counter[str]
    messages: list[str]


@dataclass(frozen=True)
class WorkOrderRow:
    handler: str
    ticket: str
    acceptance_number: str
    due_at: str


@dataclass(frozen=True)
class MentionResolution:
    mentionable: set[str]
    missing: set[str]


def load_config(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig") as handle:
        return json.load(handle)


def normalize_header(value: Any) -> str:
    return "" if value is None else "".join(str(value).strip().split())


def normalize_text(value: Any) -> str:
    return "" if value is None else " ".join(str(value).strip().split())


def normalize_identifier(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return normalize_text(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_text(value)


def file_digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def messages_digest(messages: list[str]) -> str:
    digest = hashlib.sha256()
    for message in messages:
        digest.update(message.encode("utf-8"))
        digest.update(b"\0")
    return digest.hexdigest()


def is_supported_input(path: Path) -> bool:
    return (
        path.is_file()
        and path.suffix.lower() in SUPPORTED_SUFFIXES
        and not path.name.startswith("~$")
        and not path.name.startswith(".")
    )


def read_csv_rows(path: Path) -> tuple[list[str], Iterable[tuple[Any, ...]]]:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            handle = path.open("r", encoding=encoding, newline="")
            reader = csv.reader(handle)
            headers = next(reader)
            return headers, _closing_iterator(handle, reader)
        except UnicodeDecodeError:
            try:
                handle.close()
            except UnboundLocalError:
                pass
    raise ValueError("CSV encoding is neither UTF-8 nor GB18030")


def _closing_iterator(handle: Any, rows: Iterable[list[str]]) -> Iterable[tuple[Any, ...]]:
    try:
        for row in rows:
            yield tuple(row)
    finally:
        handle.close()


def read_excel_rows(
    path: Path, sheet_name: str | None, header_row: int
) -> tuple[list[Any], Iterable[tuple[Any, ...]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Missing dependency: install openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ValueError(f"Worksheet not found: {sheet_name}")
        sheet = workbook[sheet_name]
    else:
        sheet = workbook[workbook.sheetnames[0]]

    rows = sheet.iter_rows(values_only=True)
    for _ in range(max(0, header_row - 1)):
        next(rows, None)
    headers = list(next(rows, ()))

    def iterator() -> Iterable[tuple[Any, ...]]:
        try:
            yield from rows
        finally:
            workbook.close()

    return headers, iterator()


def read_rows(path: Path, excel_config: dict[str, Any]) -> tuple[list[Any], Iterable[tuple[Any, ...]]]:
    if path.suffix.lower() == ".csv":
        return read_csv_rows(path)
    return read_excel_rows(path, excel_config.get("sheet"), int(excel_config.get("header_row", 1)))


def resolve_column(headers: list[Any], candidates: list[str], required: bool) -> int | None:
    normalized = {normalize_header(header): index for index, header in enumerate(headers)}
    for candidate in candidates:
        index = normalized.get(normalize_header(candidate))
        if index is not None:
            return index
    if required:
        raise ValueError(f"Missing required column; expected one of: {', '.join(candidates)}")
    return None


def value_at(row: tuple[Any, ...], index: int | None) -> Any:
    return None if index is None or index >= len(row) else row[index]


def collect_work_orders(path: Path, config: dict[str, Any]) -> tuple[int, list[WorkOrderRow]]:
    excel_config = config["excel"]
    headers, rows = read_rows(path, excel_config)
    strict_columns = bool(excel_config.get("strict_required_columns", False))
    handler_index = resolve_column(headers, excel_config["handler_columns"], required=True)
    status_index = resolve_column(headers, excel_config["status_columns"], required=False)
    ticket_index = resolve_column(headers, excel_config["ticket_columns"], required=strict_columns)
    acceptance_index = resolve_column(
        headers,
        excel_config.get("acceptance_columns", ["受理号码"]),
        required=strict_columns,
    )
    due_index = resolve_column(
        headers,
        excel_config.get("due_columns", ["本环节到期时限"]),
        required=strict_columns,
    )
    excluded = {normalize_text(value).casefold() for value in excel_config.get("excluded_statuses", [])}
    aliases = config["message"].get("name_aliases", {})
    strict_empty_handler = bool(excel_config.get("strict_empty_handler", False))

    source_rows = list(rows)
    handler_by_row: dict[int, str] = {}
    workgroup_by_row: dict[int, str] = {}
    workgroup_index = resolve_column(
        headers,
        excel_config.get("workgroup_columns", ["当前处理工作组"]),
        required=strict_columns,
    )
    first_data_row = int(excel_config.get("header_row", 1)) + 1
    for row_number, row in enumerate(source_rows, start=first_data_row):
        handler_by_row[row_number] = normalize_text(value_at(row, handler_index))
        workgroup_by_row[row_number] = normalize_text(value_at(row, workgroup_index))

    def infer_handler(row_number: int) -> str:
        workgroup = workgroup_by_row.get(row_number, "")
        if not workgroup:
            return ""
        candidates = [
            (abs(candidate_row - row_number), candidate_row, candidate_handler)
            for candidate_row, candidate_handler in handler_by_row.items()
            if candidate_handler and workgroup_by_row.get(candidate_row) == workgroup
        ]
        if not candidates:
            return ""
        candidates.sort(key=lambda item: (item[0], item[1]))
        nearest_distance = candidates[0][0]
        nearest_names = {name for distance, _, name in candidates if distance == nearest_distance}
        if len(nearest_names) != 1:
            return ""
        return candidates[0][2]

    total_rows = 0
    work_orders: list[WorkOrderRow] = []
    for row_number, row in enumerate(source_rows, start=first_data_row):
        if not any(value not in (None, "") for value in row):
            continue
        total_rows += 1
        handler = normalize_text(value_at(row, handler_index))
        if not handler:
            handler = infer_handler(row_number)
            if handler:
                print(
                    f"inferred 当前处理人 at row {row_number}: {handler} "
                    f"(当前处理工作组={workgroup_by_row.get(row_number)!r})"
                )
        if not handler:
            if strict_empty_handler:
                raise ValueError(f"Row {row_number}: 当前处理人为空")
            logging.warning("Row %s has no resolvable 当前处理人; skipped", row_number)
            continue
        status = normalize_text(value_at(row, status_index)).casefold()
        if status and status in excluded:
            continue
        handler = normalize_text(aliases.get(handler, handler))
        if handler:
            work_orders.append(
                WorkOrderRow(
                    handler=handler,
                    ticket=normalize_identifier(value_at(row, ticket_index)),
                    acceptance_number=normalize_identifier(value_at(row, acceptance_index)),
                    due_at=normalize_text(value_at(row, due_index)),
                )
            )
    return total_rows, work_orders


def collect_pending(path: Path, config: dict[str, Any]) -> tuple[int, int, Counter[str]]:
    total_rows, work_orders = collect_work_orders(path, config)
    return total_rows, len(work_orders), Counter(row.handler for row in work_orders)


def mention_label(name: str, mentionable_names: set[str] | None = None) -> str:
    if mentionable_names is None or name in mentionable_names:
        return f"@{name}"
    return name


def append_missing_member_note(body: str, names: Iterable[str], mentionable_names: set[str] | None = None) -> str:
    if mentionable_names is None:
        return body
    missing = sorted(set(names) - mentionable_names)
    if not missing:
        return body
    return body + "\n\n" + "、".join(missing) + " 没查到，请问是否在本群里面。"


def classify_file(path: Path, config: dict[str, Any]) -> str:
    name = path.stem.casefold()
    modes = config.get("file_modes", {})
    if any(normalize_text(keyword).casefold() in name for keyword in modes.get("overdue_keywords", ["过期", "超期"])):
        return "overdue"
    if any(normalize_text(keyword).casefold() in name for keyword in modes.get("detail_keywords", ["投诉", "当天到期", "今日到期"])):
        return "detail"
    return "summary"


def add_batch_labels(messages: list[str], title: str) -> list[str]:
    total = len(messages)
    return [
        f"【{title}{f' {index}/{total}' if total > 1 else ''}】\n\n{body}"
        for index, body in enumerate(messages, start=1)
    ]


def build_overdue_messages(
    work_orders: list[WorkOrderRow],
    config: dict[str, Any],
    mentionable_names: set[str] | None = None,
) -> list[str]:
    counts = Counter(row.handler for row in work_orders)
    if not counts:
        return []
    message_config = config["message"]
    max_people = max(1, int(message_config.get("max_people_per_message", 8)))
    max_characters = max(200, int(message_config.get("max_characters", 1800)))
    title = normalize_text(message_config.get("overdue_title", "已过期工单催办"))
    entries = [
        (name, f"{mention_label(name, mentionable_names)} 还有 {count} 单已过期未处理，请尽快处理。")
        for name, count in sorted(counts.items())
    ]
    bodies: list[str] = []
    current: list[tuple[str, str]] = []
    for name, line in entries:
        candidate = current + [(name, line)]
        if current and (
            len(candidate) > max_people
            or len("\n".join(line for _, line in candidate)) + len(title) + 20 > max_characters
        ):
            bodies.append(
                append_missing_member_note(
                    "\n".join(line for _, line in current),
                    [name for name, _ in current],
                    mentionable_names,
                )
            )
            current = [(name, line)]
        else:
            current = candidate
    if current:
        bodies.append(
            append_missing_member_note(
                "\n".join(line for _, line in current),
                [name for name, _ in current],
                mentionable_names,
            )
        )
    return add_batch_labels(bodies, title)


def build_detail_messages(
    work_orders: list[WorkOrderRow],
    config: dict[str, Any],
    mentionable_names: set[str] | None = None,
) -> list[str]:
    if not work_orders:
        return []
    message_config = config["message"]
    max_characters = max(300, int(message_config.get("max_characters", 1800)))
    max_people = max(1, int(message_config.get("max_people_per_message", 8)))
    title = normalize_text(message_config.get("detail_title", "今日到期工单提醒"))
    grouped: dict[str, list[WorkOrderRow]] = {}
    for row in work_orders:
        grouped.setdefault(row.handler, []).append(row)

    blocks: list[tuple[str, str]] = []
    for name in sorted(grouped):
        rows = sorted(grouped[name], key=lambda item: (item.due_at, item.ticket))
        detail_lines = [
            f"{index}. 工单：{row.ticket or '未提供'}｜受理：{row.acceptance_number or '未提供'}"
            for index, row in enumerate(rows, start=1)
        ]
        header = f"{mention_label(name, mentionable_names)} 今日到期 {len(rows)} 单："
        current: list[str] = []
        for line in detail_lines:
            candidate = current + [line]
            estimated = len(title) + len(header) + len("\n".join(candidate)) + 30
            if current and estimated > max_characters:
                blocks.append((name, header + "\n" + "\n".join(current)))
                current = [line]
            else:
                current = candidate
        if current:
            blocks.append((name, header + "\n" + "\n".join(current)))

    bodies: list[str] = []
    current_blocks: list[str] = []
    current_people: set[str] = set()
    for name, block in blocks:
        candidate = current_blocks + [block]
        candidate_people = current_people | {name}
        if current_blocks and (
            len(candidate_people) > max_people
            or len("\n\n".join(candidate)) + len(title) + 20 > max_characters
        ):
            bodies.append(append_missing_member_note("\n\n".join(current_blocks), current_people, mentionable_names))
            current_blocks = [block]
            current_people = {name}
        else:
            current_blocks = candidate
            current_people = candidate_people
    if current_blocks:
        bodies.append(append_missing_member_note("\n\n".join(current_blocks), current_people, mentionable_names))
    return add_batch_labels(bodies, title)


def build_messages_for_file(
    path: Path,
    work_orders: list[WorkOrderRow],
    config: dict[str, Any],
    mentionable_names: set[str] | None = None,
) -> list[str]:
    mode = classify_file(path, config)
    if mode == "detail":
        return build_detail_messages(work_orders, config, mentionable_names)
    if mode == "overdue":
        return build_overdue_messages(work_orders, config, mentionable_names)
    raise ValueError(
        f"Unsupported work-order file name: {path.name}. "
        "File name must contain 投诉/当天到期/今日到期 or 过期/超期."
    )


def validate_delivery_config(config: dict[str, Any]) -> None:
    delivery = config.get("delivery", {})
    target = normalize_text(delivery.get("target"))
    if not re.fullmatch(r"group:[A-Za-z0-9_-]+", target):
        raise ValueError(f"Invalid Yuanbao group target: {target!r}. Expected format: group:<groupCode>.")
    openclaw_cmd = normalize_text(delivery.get("openclaw_cmd"))
    if not openclaw_cmd or not Path(openclaw_cmd).exists():
        raise ValueError(f"OpenClaw command not found: {openclaw_cmd!r}")


def ensure_input_file_ready(path: Path, config: dict[str, Any]) -> None:
    try:
        first_size = path.stat().st_size
        with path.open("rb") as handle:
            handle.read(1)
    except OSError as exc:
        raise ValueError(f"Input file is not readable or may be occupied: {path.name}") from exc
    stable_seconds = float(config.get("input_safety", {}).get("stable_seconds", 0))
    if stable_seconds <= 0:
        return
    time.sleep(stable_seconds)
    second_size = path.stat().st_size
    if first_size != second_size:
        raise ValueError(f"Input file is still changing, wait until copy/export is complete: {path.name}")


def validate_work_order_values(path: Path, work_orders: list[WorkOrderRow], config: dict[str, Any]) -> None:
    excel_config = config.get("excel", {})
    if not excel_config.get("strict_field_values", False):
        return
    acceptance_pattern = excel_config.get("acceptance_number_pattern")
    errors: list[str] = []
    for index, row in enumerate(work_orders, start=1):
        if not row.handler:
            errors.append(f"record {index}: 当前处理人为空")
        if not row.ticket:
            errors.append(f"record {index}: 工单流水号为空")
        if not row.acceptance_number:
            errors.append(f"record {index}: 受理号码为空")
        for label, value in (("工单流水号", row.ticket), ("受理号码", row.acceptance_number)):
            if SCIENTIFIC_NOTATION_RE.fullmatch(value):
                errors.append(f"record {index}: {label}疑似科学计数法 {value!r}")
        if acceptance_pattern and row.acceptance_number and not re.fullmatch(acceptance_pattern, row.acceptance_number):
            errors.append(f"record {index}: 受理号码格式异常 {row.acceptance_number!r}")
    if errors:
        preview = "; ".join(errors[:8])
        suffix = f"; ... 共 {len(errors)} 项" if len(errors) > 8 else ""
        raise ValueError(f"{path.name} failed value validation: {preview}{suffix}")


def validate_message_limits(messages: list[str], config: dict[str, Any]) -> None:
    message_config = config.get("message", {})
    max_characters = int(message_config.get("max_characters", 1800))
    max_mentions = int(message_config.get("max_mentions_per_message", message_config.get("max_people_per_message", 8)))
    for index, message in enumerate(messages, start=1):
        if len(message) > max_characters:
            raise ValueError(
                f"Message {index} exceeds max_characters after splitting: {len(message)} > {max_characters}"
            )
        mention_count = len(set(MENTION_RE.findall(message)))
        if mention_count > max_mentions:
            raise ValueError(
                f"Message {index} exceeds max_mentions_per_message after splitting: {mention_count} > {max_mentions}"
            )


def ensure_gateway_ready(config: dict[str, Any]) -> None:
    delivery = config["delivery"]
    if not delivery.get("preflight_gateway", True):
        return
    command = [delivery["openclaw_cmd"], "gateway", "status"]
    timeout = max(1, int(delivery.get("gateway_ready_timeout_seconds", 60)))
    interval = max(0.5, float(delivery.get("gateway_ready_poll_seconds", 2)))
    deadline = time.monotonic() + timeout
    required_markers = delivery.get(
        "gateway_ready_markers",
        ["Connectivity probe: ok", "Listening:"],
    )
    last_output = ""
    while time.monotonic() < deadline:
        completed = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace")
        last_output = (completed.stdout or "") + "\n" + (completed.stderr or "")
        if completed.returncode == 0 and all(marker in last_output for marker in required_markers):
            return
        time.sleep(interval)
    missing = [marker for marker in required_markers if marker not in last_output]
    raise RuntimeError(
        f"Gateway was not ready within {timeout}s; missing status marker(s): "
        + ", ".join(missing)
        + "\n"
        + last_output.strip()
    )


def init_db(path: Path) -> sqlite3.Connection:
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path)
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS imports (
            digest TEXT PRIMARY KEY,
            source_name TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            total_rows INTEGER NOT NULL,
            pending_rows INTEGER NOT NULL,
            message_count INTEGER NOT NULL,
            status TEXT NOT NULL,
            error TEXT
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS deliveries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            digest TEXT NOT NULL,
            batch_number INTEGER NOT NULL,
            message TEXT NOT NULL,
            cron_job_id TEXT,
            status TEXT NOT NULL,
            created_at TEXT NOT NULL,
            error TEXT,
            UNIQUE(digest, batch_number)
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS send_batches (
            id TEXT PRIMARY KEY,
            file_digest TEXT NOT NULL,
            target TEXT NOT NULL,
            message_digest TEXT NOT NULL,
            source_name TEXT NOT NULL,
            created_at TEXT NOT NULL,
            status TEXT NOT NULL,
            error TEXT
        )
        """
    )
    connection.commit()
    return connection


def already_imported(connection: sqlite3.Connection, digest: str) -> bool:
    return connection.execute("SELECT 1 FROM imports WHERE digest = ?", (digest,)).fetchone() is not None


def reserve_send_batch(
    connection: sqlite3.Connection,
    file_digest_value: str,
    target: str,
    message_digest_value: str,
    source_name: str,
) -> str | None:
    try:
        connection.execute("BEGIN IMMEDIATE")
        running = connection.execute(
            """
            SELECT id FROM send_batches
            WHERE file_digest = ? AND target = ? AND message_digest = ? AND status = 'sending'
            """,
            (file_digest_value, target, message_digest_value),
        ).fetchone()
        if running:
            connection.rollback()
            return None
        batch_id = hashlib.sha256(
            f"{file_digest_value}\0{target}\0{message_digest_value}\0{datetime.now().isoformat()}".encode("utf-8")
        ).hexdigest()[:16]
        now = datetime.now().astimezone().isoformat(timespec="seconds")
        connection.execute(
            "INSERT INTO send_batches VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (batch_id, file_digest_value, target, message_digest_value, source_name, now, "sending", None),
        )
        connection.commit()
        return batch_id
    except Exception:
        connection.rollback()
        raise


def update_send_batch(
    connection: sqlite3.Connection,
    batch_id: str,
    status: str,
    error: str | None = None,
) -> None:
    connection.execute(
        "UPDATE send_batches SET status = ?, error = ? WHERE id = ?",
        (status, error, batch_id),
    )
    connection.commit()


def build_cron_prompt(message: str) -> str:
    marker = "[[YB_LINE_BREAK]]"
    if marker in message:
        raise ValueError(f"Message contains reserved line-break marker: {marker}")
    encoded = message.replace("\r\n", "\n").replace("\r", "\n").replace("\n", marker)
    return (
        f"将下面文本中的每个{marker}替换为一个真实换行，然后原样输出。"
        "不要解释、改写、增删或调用工具，只输出还原后的催办消息：" + encoded
    )


def create_cron(message: str, batch_number: int, config: dict[str, Any]) -> str:
    delivery = config["delivery"]
    prompt = build_cron_prompt(message)
    command = [
        delivery["openclaw_cmd"],
        "cron",
        "add",
        "--name",
        f"work-order-reminder-{datetime.now():%Y%m%d-%H%M%S}-{batch_number}",
        "--at",
        delivery.get("delay", "+30s"),
        "--session",
        "isolated",
        "--light-context",
        "--model",
        delivery.get("model", "deepseek/deepseek-chat"),
        "--thinking",
        "off",
        "--announce",
        "--channel",
        delivery.get("channel", "yuanbao"),
        "--account",
        delivery.get("account", "default"),
        "--to",
        delivery["target"],
        "--delete-after-run",
        "--json",
        "--message",
        prompt,
    ]
    completed = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if completed.returncode != 0:
        raise RuntimeError((completed.stderr or completed.stdout).strip())
    start = completed.stdout.find("{")
    if start < 0:
        raise RuntimeError(f"OpenClaw returned no JSON: {completed.stdout.strip()}")
    payload = json.loads(completed.stdout[start:])
    saved_delivery = payload.get("delivery", {})
    expected_channel = delivery.get("channel", "yuanbao")
    expected_target = delivery["target"]
    if saved_delivery.get("channel") != expected_channel or saved_delivery.get("to") != expected_target:
        raise RuntimeError(
            "OpenClaw did not persist the explicit delivery target: "
            f"expected {expected_channel}:{expected_target}, got {saved_delivery}"
        )
    job_id = payload.get("id")
    if not job_id:
        raise RuntimeError(f"OpenClaw response has no job id: {payload}")
    return str(job_id)


def parse_iso_ms(value: str | None) -> int | None:
    if not value:
        return None
    try:
        normalized = value.replace("Z", "+00:00")
        return int(datetime.fromisoformat(normalized).timestamp() * 1000)
    except ValueError:
        return None


def yuanbao_send_ok_from_logs(entry: dict[str, Any], config: dict[str, Any]) -> str | None:
    delivery = config["delivery"]
    group_code = delivery["target"].removeprefix("group:")
    log_dir = Path(delivery.get("gateway_log_dir", Path(tempfile.gettempdir()) / "openclaw"))
    if not log_dir.exists():
        return None
    start_ms = int(entry.get("runAtMs") or parse_iso_ms(entry.get("runAtIso")) or 0) - 1000
    end_ms = int(entry.get("ts") or parse_iso_ms(entry.get("tsIso")) or int(time.time() * 1000)) + 5000
    log_files = sorted(log_dir.glob("openclaw-*.log"), key=lambda path: path.stat().st_mtime, reverse=True)[:5]
    for log_path in log_files:
        try:
            with log_path.open("r", encoding="utf-8", errors="replace") as handle:
                for line in handle:
                    if "[group] send ok" not in line or group_code not in line:
                        continue
                    try:
                        record = json.loads(line)
                    except json.JSONDecodeError:
                        continue
                    event_ms = parse_iso_ms(record.get("_meta", {}).get("time") or record.get("time"))
                    if event_ms is not None and not (start_ms <= event_ms <= end_ms):
                        continue
                    message = normalize_text(record.get("message", ""))
                    match = re.search(r'"msgId":"([^"]+)"', message)
                    return match.group(1) if match else "send-ok"
        except OSError as exc:
            logging.warning("Could not scan Yuanbao delivery log %s: %s", log_path, exc)
    return None


def wait_for_cron(job_id: str, config: dict[str, Any]) -> dict[str, Any]:
    delivery = config["delivery"]
    timeout = max(15, int(delivery.get("wait_timeout_seconds", 180)))
    deadline = time.monotonic() + timeout
    command = [
        delivery["openclaw_cmd"],
        "cron",
        "runs",
        "--id",
        job_id,
        "--limit",
        "1",
    ]
    while time.monotonic() < deadline:
        completed = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace")
        if completed.returncode == 0:
            start = completed.stdout.find("{")
            if start >= 0:
                payload = json.loads(completed.stdout[start:])
                entries = payload.get("entries", [])
                if entries:
                    entry = entries[0]
                    if entry.get("status") != "ok":
                        raise RuntimeError(entry.get("error") or f"Cron job failed: {entry}")
                    if entry.get("deliveryStatus") != "delivered":
                        msg_id = yuanbao_send_ok_from_logs(entry, config)
                        if msg_id:
                            entry["deliveryStatus"] = "delivered-via-yuanbao-log"
                            entry["yuanbaoMsgId"] = msg_id
                            print(f"confirmed Yuanbao send ok from plugin log: {msg_id}")
                            return entry
                        if not delivery.get("allow_undelivered_status", False):
                            raise RuntimeError(
                                "OpenClaw cron finished but deliveryStatus is not delivered: "
                                + json.dumps(entry, ensure_ascii=False)
                            )
                        print(
                            "warning: OpenClaw cron did not mark delivery as delivered; "
                            "the Yuanbao plugin/server response remains authoritative."
                        )
                    return entry
        time.sleep(2)
    raise TimeoutError(f"Timed out waiting {timeout}s for cron job {job_id}")


def iter_string_values(value: Any) -> Iterable[str]:
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for item in value.values():
            yield from iter_string_values(item)
    elif isinstance(value, list):
        for item in value:
            yield from iter_string_values(item)


def extract_inbound_member(log_line: str) -> tuple[str, str, str] | None:
    try:
        outer = json.loads(log_line)
    except json.JSONDecodeError:
        return None
    marker = "decoded message "
    decoder = json.JSONDecoder()
    for text in iter_string_values(outer):
        marker_index = text.find(marker)
        if marker_index < 0:
            continue
        payload_text = text[marker_index + len(marker):].lstrip()
        try:
            payload, _ = decoder.raw_decode(payload_text)
        except json.JSONDecodeError:
            continue
        if payload.get("chatType") != "group":
            continue
        group_code = normalize_text(payload.get("groupCode"))
        nickname = normalize_text(payload.get("senderNickname"))
        user_id = normalize_text(payload.get("fromAccount"))
        if group_code and nickname and user_id and not user_id.startswith("bot_"):
            return group_code, nickname, user_id
    return None


def sync_member_mapping_from_logs(config: dict[str, Any]) -> list[str]:
    delivery = config["delivery"]
    member_file = Path(delivery["member_mapping_file"])
    group_code = delivery["target"].removeprefix("group:")
    log_dir = Path(delivery.get("gateway_log_dir", Path(tempfile.gettempdir()) / "openclaw"))
    if member_file.exists():
        with member_file.open("r", encoding="utf-8-sig") as handle:
            mapping = json.load(handle)
    else:
        mapping = {}
    members = mapping.setdefault(group_code, {})
    learned: list[str] = []
    if log_dir.exists():
        log_files = sorted(log_dir.glob("openclaw-*.log"), key=lambda path: path.stat().st_mtime, reverse=True)
        for log_path in log_files:
            try:
                with log_path.open("r", encoding="utf-8", errors="replace") as handle:
                    for line in handle:
                        record = extract_inbound_member(line)
                        if not record:
                            continue
                        record_group, nickname, user_id = record
                        if record_group == group_code and members.get(nickname) != user_id:
                            members[nickname] = user_id
                            learned.append(nickname)
            except OSError as exc:
                logging.warning("Could not scan Yuanbao log %s: %s", log_path, exc)
    if learned:
        member_file.parent.mkdir(parents=True, exist_ok=True)
        temporary = member_file.with_suffix(member_file.suffix + ".tmp")
        with temporary.open("w", encoding="utf-8", newline="\n") as handle:
            json.dump(mapping, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        temporary.replace(member_file)
    return sorted(set(learned))


def resolve_mentions(counts: Counter[str], config: dict[str, Any]) -> MentionResolution:
    learned = sync_member_mapping_from_logs(config)
    if learned:
        print(f"learned member mapping: {', '.join(learned)}")
    member_file = Path(config["delivery"]["member_mapping_file"])
    if member_file.exists():
        with member_file.open("r", encoding="utf-8-sig") as handle:
            mapping = json.load(handle)
    else:
        mapping = {}
    group_code = config["delivery"]["target"].removeprefix("group:")
    members = mapping.get(group_code, {})
    ambiguous = sorted(set(config.get("member_safety", {}).get("ambiguous_names", [])) & set(counts))
    if ambiguous:
        print(
            f"warning: ambiguous Yuanbao member nickname(s) in group {group_code}, will not native @: "
            + ", ".join(ambiguous)
        )
    missing = {name for name in counts if not members.get(name)} | set(ambiguous)
    if missing:
        print(
            f"warning: Yuanbao userId not resolved in group {group_code}; will send without native @: "
            + ", ".join(sorted(missing))
        )
    mentionable = {name for name in counts if members.get(name) and name not in missing}
    return MentionResolution(mentionable=mentionable, missing=missing)


def record_import(
    connection: sqlite3.Connection,
    result: ImportResult,
    status: str,
    deliveries: list[tuple[int, str, str | None, str, str | None]],
    error: str | None = None,
) -> None:
    now = datetime.now().astimezone().isoformat(timespec="seconds")
    connection.execute(
        "INSERT OR REPLACE INTO imports VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        (result.digest, result.source.name, now, result.total_rows, result.pending_rows, len(result.messages), status, error),
    )
    for batch_number, message, job_id, delivery_status, delivery_error in deliveries:
        connection.execute(
            "INSERT OR REPLACE INTO deliveries (digest, batch_number, message, cron_job_id, status, created_at, error) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            (result.digest, batch_number, message, job_id, delivery_status, now, delivery_error),
        )
    connection.commit()


def move_with_timestamp(source: Path, destination: Path) -> Path:
    destination.mkdir(parents=True, exist_ok=True)
    target = destination / source.name
    if target.exists():
        target = destination / f"{source.stem}_{datetime.now():%Y%m%d_%H%M%S}{source.suffix}"
    return Path(shutil.move(str(source), str(target)))


def process_file(path: Path, config: dict[str, Any], connection: sqlite3.Connection, dry_run: bool) -> str:
    ensure_input_file_ready(path, config)
    digest = file_digest(path)
    total_rows, work_orders = collect_work_orders(path, config)
    pending_rows = len(work_orders)
    counts = Counter(row.handler for row in work_orders)
    validate_work_order_values(path, work_orders, config)
    mode = classify_file(path, config)
    if mode == "summary":
        raise ValueError(
            f"Unsupported work-order file name: {path.name}. "
            "File name must contain 投诉/当天到期/今日到期 or 过期/超期."
        )
    mentionable_names: set[str] | None = None
    should_try_delivery = bool(work_orders and not dry_run and config["delivery"].get("enabled", False))
    if should_try_delivery:
        validate_delivery_config(config)
        ensure_gateway_ready(config)
        mention_resolution = resolve_mentions(counts, config)
        mentionable_names = mention_resolution.mentionable
    messages = build_messages_for_file(path, work_orders, config, mentionable_names)
    validate_message_limits(messages, config)
    should_deliver = bool(messages and not dry_run and config["delivery"].get("enabled", False))
    batch_id: str | None = None
    if should_deliver:
        batch_id = reserve_send_batch(
            connection,
            digest,
            config["delivery"]["target"],
            messages_digest(messages),
            path.name,
        )
        if not batch_id:
            raise RuntimeError(
                f"Same file/message/target is already sending: {path.name}. "
                "Wait for the current run to finish before clicking send again."
            )
    result = ImportResult(path, digest, total_rows, pending_rows, counts, messages)
    deliveries: list[tuple[int, str, str | None, str, str | None]] = []

    try:
        print(f"\n=== {path.name} ===")
        print(f"rows={total_rows}, pending={pending_rows}, people={len(counts)}, messages={len(messages)}")
        for batch_number, message in enumerate(messages, start=1):
            print(f"\n--- message {batch_number}/{len(messages)} ---\n{message}")
            if not should_deliver:
                deliveries.append((batch_number, message, None, "preview", None))
                continue
            try:
                job_id = create_cron(message, batch_number, config)
                print(f"scheduled cron job: {job_id}")
                entry = wait_for_cron(job_id, config)
                deliveries.append((batch_number, message, job_id, "delivered", None))
                print(f"delivered: {entry.get('tsIso', entry.get('ts'))}")
            except Exception as exc:
                deliveries.append((batch_number, message, None, "error", str(exc)))
                raise

        status = "no_pending" if not messages else ("preview" if not should_deliver else "delivered")
        record_import(connection, result, status, deliveries)
        if batch_id:
            update_send_batch(connection, batch_id, status)
        return status
    except Exception as exc:
        if batch_id:
            update_send_batch(connection, batch_id, "error", str(exc))
        raise


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Group pending work orders by handler and schedule YuanBao reminders.")
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--file", type=Path, help="Process one Excel/CSV file instead of scanning inbox.")
    parser.add_argument("--send", action="store_true", help="Enable delivery for this run (config must also enable it).")
    parser.add_argument("--keep", action="store_true", help="Do not move processed inbox files.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    config = load_config(args.config.resolve())
    paths = {key: Path(value) for key, value in config["paths"].items()}
    for key in ("inbox", "archive", "failed"):
        paths[key].mkdir(parents=True, exist_ok=True)
    connection = init_db(paths["state_db"])
    dry_run = not (args.send and config["delivery"].get("enabled", False))
    if args.file:
        explicit_file = args.file.resolve()
        if not is_supported_input(explicit_file):
            print(f"Unsupported or temporary input file: {explicit_file}", file=sys.stderr)
            return 1
        candidates = [explicit_file]
    else:
        candidates = sorted(path for path in paths["inbox"].iterdir() if is_supported_input(path))
    if not candidates:
        print("No supported Excel/CSV files found.")
        return 0

    failures = 0
    for path in candidates:
        try:
            status = process_file(path, config, connection, dry_run)
            print(status)
            if not dry_run and not args.file and not args.keep and not status.startswith("SKIP"):
                move_with_timestamp(path, paths["archive"])
        except Exception as exc:
            failures += 1
            logging.exception("Failed to process %s", path)
            if not dry_run and not args.file and not args.keep:
                try:
                    move_with_timestamp(path, paths["failed"])
                except OSError as move_error:
                    logging.error("Could not move failed input %s: %s", path, move_error)
            print(f"ERROR {path.name}: {exc}", file=sys.stderr)
    connection.close()
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())

